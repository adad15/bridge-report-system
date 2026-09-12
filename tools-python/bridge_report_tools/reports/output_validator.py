"""最终 DOCX 校验（设计 §20）。

刷完域之后、交付之前的最后一道关。跟模板校验（template_validator）不是一回事：
那边查的是"这份骨架能不能用来生成"，这边查的是"这份成品能不能交出去"。

一条纪律：**宁可不交付，也不交一份看着正常、实际残缺的报告。** 锚点没替换掉、
占位符还剩着、图片关系断了、分节没了——这些在 Word 里打开都不报错，读者却会拿到
一份错的报告。所以每一条都在这里查死。
"""

from __future__ import annotations

import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from bridge_report_tools.reports.contract import (
    ANCHOR_PREFIX,
    FIELDS_ALLOWED_ANYWHERE,
    FIELDS_ALLOWED_INSIDE_TOC,
    ReportContract,
)
from bridge_report_tools.reports.docx_scan import scan_document
from bridge_report_tools.reports.errors import ReportTemplateError
from bridge_report_tools.reports.issues import TemplateIssue, error

CONTENT_TYPES_ENTRY = "[Content_Types].xml"
DOCUMENT_ENTRY = "word/document.xml"

#: 图片引用。r:embed 是内嵌图，r:link 是外链图——外链图在这里一律算断（§20 第 5 条）。
EMBED_PATTERN = re.compile(r'r:embed="([^"]+)"')
LINK_PATTERN = re.compile(r'r:link="([^"]+)"')


@dataclass(frozen=True)
class OutputValidationResult:
    issues: list[TemplateIssue]
    section_count: int = 0
    image_count: int = 0

    @property
    def status(self) -> str:
        return "invalid" if any(i.severity == "error" for i in self.issues) else "valid"


def _check_package(path: Path) -> list[TemplateIssue]:
    """第 1、8 条：是结构完整的 OOXML ZIP，且能被解析器再打开一次。"""
    try:
        with zipfile.ZipFile(path) as archive:
            broken = archive.testzip()
            if broken is not None:
                return [
                    error("report_output_not_ooxml", f"生成的文件里 {broken} 这一项已损坏。")
                ]
            names = set(archive.namelist())
    except zipfile.BadZipFile as exc:
        return [error("report_output_not_ooxml", f"生成的文件不是有效的 docx：{exc}")]

    missing = [entry for entry in (CONTENT_TYPES_ENTRY, DOCUMENT_ENTRY) if entry not in names]
    if missing:
        return [
            error("report_output_not_ooxml", f"生成的文件缺少必要条目：{'、'.join(missing)}。")
        ]
    return []


def _check_leftovers(document, contract: ReportContract) -> list[TemplateIssue]:
    """第 2、3 条：不残留锚点，不残留未知占位符。

    锚点直接按文本找，不走扫描器：扫描器只认合法的 `[[REPORT:块]]`，而装配失败留下
    的半截标记（比如只剩 `[[REPORT:`）同样不能交付。
    """
    issues: list[TemplateIssue] = []
    scan = scan_document(document)

    for location, text in _iter_texts(document):
        if ANCHOR_PREFIX in text:
            issues.append(
                error(
                    "report_output_anchor_left",
                    f"成品里还留着未装配的内容锚点：{text.strip()[:60]}",
                    location,
                )
            )

    for hit in scan.placeholders:
        issues.append(
            error(
                "report_output_placeholder_left",
                f"成品里还留着未替换的占位符 {{{{{hit.name}}}}}。",
                hit.location,
            )
        )
    for location in scan.unbalanced_braces:
        issues.append(
            error(
                "report_output_placeholder_left",
                "成品里有配不成对的花括号，可能是被截断的占位符。",
                location,
            )
        )
    return issues


def _iter_texts(document):
    """正文、表格单元格和页眉页脚里的段落文字，带位置串。"""
    for index, paragraph in enumerate(document.paragraphs):
        yield f"body#{index}", paragraph.text
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            for cell_index, cell in enumerate(row.cells):
                yield f"table{table_index}#{row_index},{cell_index}", cell.text
    for section_index, section in enumerate(document.sections):
        for name, part in (("header", section.header), ("footer", section.footer)):
            for index, paragraph in enumerate(part.paragraphs):
                yield f"{name}:{section_index}#{index}", paragraph.text


def _check_blocks(
    contract: ReportContract, blocks_rendered: list[str]
) -> list[TemplateIssue]:
    """第 4 条：契约要求的内容块都装配过了。

    装配层报告哪些块渲染完了，这里只做对账——成品文件本身看不出"这段是哪个块出的"。
    对账放在这里而不是装配层，是因为交付前的最后一道关必须独立说一次"齐了"。
    """
    rendered = {block.partition(":")[0] for block in blocks_rendered}
    missing = sorted(contract.required_anchors - rendered)
    if not missing:
        return []
    return [
        error(
            "report_output_block_missing",
            f"契约要求的内容块没有装配：{'、'.join(missing)}。",
        )
    ]


def _check_media(path: Path, document) -> tuple[list[TemplateIssue], int]:
    """第 5 条：图片关系完整，媒体文件读得出来。"""
    issues: list[TemplateIssue] = []
    with zipfile.ZipFile(path) as archive:
        sizes = {info.filename: info.file_size for info in archive.infolist()}

    part = document.part
    embedded = set(EMBED_PATTERN.findall(part.blob.decode("utf-8", "ignore")))
    linked = set(LINK_PATTERN.findall(part.blob.decode("utf-8", "ignore")))
    for relationship_id in sorted(linked):
        issues.append(
            error(
                "report_output_media_missing",
                f"图片 {relationship_id} 是外链而不是内嵌，文件发出去就看不到图了。",
            )
        )

    for relationship_id in sorted(embedded):
        relationship = part.rels.get(relationship_id)
        if relationship is None or relationship.reltype != RT.IMAGE:
            issues.append(
                error(
                    "report_output_media_missing",
                    f"正文引用了图片 {relationship_id}，但包里没有对应的图片关系。",
                )
            )
            continue
        entry = f"word/{relationship.target_ref.lstrip('/')}"
        if sizes.get(entry, 0) <= 0:
            issues.append(
                error(
                    "report_output_media_missing",
                    f"图片关系 {relationship_id} 指向 {entry}，但这个文件不在包里或是空的。",
                )
            )
    return issues, len(embedded)


def _check_fields(document) -> list[TemplateIssue]:
    """第 6 条：只留允许的域。

    目录更新后 Word 会在 TOC 结果里自动生成一堆 PAGEREF 和 HYPERLINK，那是它自己
    写的，放行；同类域出现在目录之外仍按 §7.6 拒绝。
    """
    issues: list[TemplateIssue] = []
    for hit in scan_document(document).fields:
        if hit.keyword in FIELDS_ALLOWED_ANYWHERE:
            continue
        if hit.keyword in FIELDS_ALLOWED_INSIDE_TOC and hit.inside_toc:
            continue
        issues.append(
            error(
                "report_output_field_not_allowed",
                f"成品里出现了不允许的域 {hit.keyword or '(空指令)'}"
                f"（指令：{hit.instruction}）。",
                hit.location,
            )
        )
    return issues


def _check_sections(document) -> tuple[list[TemplateIssue], int]:
    """第 7 条：分节、纸张方向和页眉页脚关系都还在。"""
    sections = document.sections
    if not sections:
        return [error("report_output_section_broken", "成品里一个分节都没有。")], 0

    issues: list[TemplateIssue] = []
    for index, section in enumerate(sections):
        if not section.page_width or not section.page_height:
            issues.append(
                error(
                    "report_output_section_broken",
                    f"第 {index + 1} 节没有纸张尺寸，页面设置丢了。",
                    f"section#{index}",
                )
            )
        for name, part in (("页眉", section.header), ("页脚", section.footer)):
            try:
                part.paragraphs  # noqa: B018 - 关系断了会在这里抛
            except Exception as exc:  # pragma: no cover - 取决于损坏方式
                issues.append(
                    error(
                        "report_output_section_broken",
                        f"第 {index + 1} 节的{name}关系读不出来：{exc}",
                        f"section#{index}",
                    )
                )
    return issues, len(sections)


def _check_location(path: Path, job_directory: Path) -> list[TemplateIssue]:
    """第 9 条：成品必须落在本次任务的临时目录里。

    下载接口按任务 ID 取路径，路径一旦能指到任务目录之外，就等于给了一条读任意文件
    的通道（§22 最后一条）。
    """
    try:
        path.resolve().relative_to(job_directory.resolve())
    except ValueError:
        return [
            error(
                "report_output_path_outside_job",
                f"成品路径 {path} 不在本次任务的临时目录 {job_directory} 内。",
            )
        ]
    return []


def validate_output(
    path: Path,
    job_directory: Path,
    contract: ReportContract,
    blocks_rendered: list[str],
) -> OutputValidationResult:
    """按设计 §20 逐条查一份成品。

    包本身坏掉时直接返回，后面每一条都要先能把文档打开。
    """
    if not path.is_file():
        raise ReportTemplateError("report_output_missing", f"生成的文件不存在：{path}")

    issues = _check_location(path, job_directory)
    package_issues = _check_package(path)
    if package_issues:
        return OutputValidationResult(issues=issues + package_issues)

    try:
        document = Document(str(path))
    except Exception as exc:
        return OutputValidationResult(
            issues=issues
            + [error("report_output_not_ooxml", f"生成的文件打不开：{exc}")]
        )

    issues += _check_leftovers(document, contract)
    issues += _check_blocks(contract, blocks_rendered)
    media_issues, image_count = _check_media(path, document)
    issues += media_issues
    issues += _check_fields(document)
    section_issues, section_count = _check_sections(document)
    issues += section_issues

    return OutputValidationResult(
        issues=issues, section_count=section_count, image_count=image_count
    )
