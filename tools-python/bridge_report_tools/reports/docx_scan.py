"""按文档顺序扫描模板 docx 里的占位符、内容锚点和 Word 域。

这一层只负责"看见什么"，不判断合不合规——判断在 template_validator 里做，因为
同一份扫描结果既要给模板校验用，也要给 Docx Builder 的替换用。
"""

from __future__ import annotations

import re
from bisect import bisect_right
from dataclasses import dataclass
from pathlib import Path
from typing import Iterator

from docx import Document
from docx.oxml.ns import qn
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from bridge_report_tools.reports.contract import ANCHOR_PREFIX, TABLE_MARKER_PREFIX


CUSTOM_XML_TAG = qn("w:customXml")
FLD_CHAR_TAG = qn("w:fldChar")
FLD_CHAR_TYPE_ATTR = qn("w:fldCharType")
INSTR_TEXT_TAG = qn("w:instrText")
FLD_SIMPLE_TAG = qn("w:fldSimple")
FLD_SIMPLE_INSTR_ATTR = qn("w:instr")

# 占位符名称限定为 ASCII 标识符加点号；写成中文或带空格的一律当未知占位符报出来，
# 而不是静默漏扫——漏扫的后果是它原样出现在交付给用户的报告里。
PLACEHOLDER_PATTERN = re.compile(r"\{\{\s*(?P<name>[^{}]*?)\s*\}\}")
ANCHOR_PATTERN = re.compile(r"\[\[REPORT:\s*(?P<name>[^\[\]]*?)\s*\]\]")
TABLE_MARKER_PATTERN = re.compile(r"\[\[TABLE:\s*(?P<name>[^\[\]]*?)\s*\]\]")
VALID_PLACEHOLDER_NAME = re.compile(r"^[A-Za-z0-9_]+(?:\.[A-Za-z0-9_]+)*$")
VALID_ANCHOR_NAME = re.compile(r"^[A-Z][A-Z0-9_]*$")

LOCATION_BODY = "body"


@dataclass(frozen=True)
class RunSpan:
    """占位符在段落里横跨的 run 区间，右开。

    Word 会把 `{{report_no}}` 拆成任意多个 run（拼写检查、语言标记、修订痕迹都会拆），
    所以替换必须按区间做：文本写回首个 run 以继承它的样式，区间内其余部分清空。
    """

    start_run: int
    start_offset: int
    end_run: int
    end_offset: int


@dataclass(frozen=True)
class PlaceholderHit:
    name: str
    raw: str
    location: str
    paragraph_index: int
    span: RunSpan
    is_well_formed: bool


@dataclass(frozen=True)
class AnchorHit:
    #: 锚点内的完整文本，如 "DEFECT_TABLES" 或 "DEFECT_TABLES:SUPERSTRUCTURE"。
    name: str
    #: 冒号前的内容块名。
    block: str
    #: 冒号后的结构部位代码；不带部位时为 None。
    part: str | None
    raw: str
    location: str
    paragraph_index: int
    #: 锚点是否独占一个段落。不独占时无法整段替换，校验会拒绝（设计 §7.3）。
    owns_paragraph: bool
    paragraph_text: str


@dataclass(frozen=True)
class FieldHit:
    keyword: str
    instruction: str
    location: str
    #: 位于某个 TOC 域的结果区内部——目录更新后由 Word/WPS 自己生成，不是模板作者写的。
    inside_toc: bool


@dataclass(frozen=True)
class TableSampleHit:
    """一张表头样表：模板作者在 Word 里画好列宽和表头，生成器照着建表。

    只带"表长什么样"，不带"表里有什么"——行数和合并仍由数据决定（设计 §7.7）。
    """

    block: str
    raw: str
    paragraph_index: int
    #: 标记是否位于段首。样表整段删除，所以标记后面可以跟一句给模板作者看的说明；
    #: 但不能写在句子中间——那样删段会把人家的话一起删掉。
    starts_paragraph: bool
    #: 标记后面是不是真的跟着一张表。
    has_table: bool
    #: 表头行的文字，按列。
    headers: list[str]
    #: 列宽（EMU）。模板作者拖出来的那个宽度。
    widths: list[int]


@dataclass(frozen=True)
class TemplateScan:
    placeholders: list[PlaceholderHit]
    anchors: list[AnchorHit]
    fields: list[FieldHit]
    #: 出现了 `{{` 或 `}}` 却配不成对的段落位置，用于报错时指人。
    unbalanced_braces: list[str]
    #: 表头样表，按内容块。模板没画样表时为空——样表是可选的（设计 §7.7）。
    table_samples: list[TableSampleHit]


def visible_runs(paragraph: Paragraph) -> list[Run]:
    """段落里所有承载可见文字的 run，含超链接和文本框内的 run。

    python-docx 的 Paragraph.runs 只取 w:p 的直接子 w:r，会漏掉 w:hyperlink 里的 run。
    模板的占位符完全可能落在超链接或文本框内，漏扫等于放过一个替换不掉的占位符。
    只取带 w:t 的 run，域指令（w:instrText）因此天然被排除在文本之外。
    """
    return [Run(element, paragraph) for element in paragraph._p.xpath(".//w:r[w:t]")]


def run_text_map(runs: list[Run]) -> tuple[str, list[int]]:
    """拼接 run 文本，并返回每个 run 起点在拼接串里的下标。"""
    starts: list[int] = []
    parts: list[str] = []
    cursor = 0
    for run in runs:
        starts.append(cursor)
        text = run.text
        parts.append(text)
        cursor += len(text)
    return "".join(parts), starts


def _locate(starts: list[int], index: int) -> tuple[int, int]:
    """把拼接串下标映射回 (run 序号, run 内偏移)。"""
    run_index = bisect_right(starts, index) - 1
    return run_index, index - starts[run_index]


def span_for(starts: list[int], start: int, end: int) -> RunSpan:
    start_run, start_offset = _locate(starts, start)
    # end 是右开下标；用 end-1 定位最后一个真正被覆盖的字符，再把偏移加回 1。
    end_run, end_offset = _locate(starts, end - 1)
    return RunSpan(
        start_run=start_run,
        start_offset=start_offset,
        end_run=end_run,
        end_offset=end_offset + 1,
    )


def iter_paragraphs(parent_element, parent) -> Iterator[Paragraph]:
    """按文档顺序产出段落，下钻表格单元格和 customXml 块。

    表格用 ./w:tr/w:tc 直接取单元格元素而不是 Table.rows[].cells：后者在合并单元格上
    会把同一个 w:tc 返回多次，扫描时会把一个占位符数成两个。
    """
    for child in parent_element.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)
        elif isinstance(child, CT_Tbl):
            table = Table(child, parent)
            for tc in child.xpath("./w:tr/w:tc"):
                yield from iter_paragraphs(tc, _Cell(tc, table))
        elif child.tag == CUSTOM_XML_TAG:
            yield from iter_paragraphs(child, parent)


def _iter_containers(document) -> Iterator[tuple[str, object, object]]:
    """产出 (位置名, 容器元素, python-docx 父对象)。

    页眉页脚同样要扫：动态页眉是被禁的（设计 §7.6），标量占位符却允许出现在那里。
    linked_to_previous 的页眉没有自己的内容，跳过以免把上一节的内容重复计一遍。
    """
    yield LOCATION_BODY, document.element.body, document
    for index, section in enumerate(document.sections, start=1):
        parts = (
            ("header", section.header),
            ("footer", section.footer),
            ("header:first", section.first_page_header),
            ("footer:first", section.first_page_footer),
            ("header:even", section.even_page_header),
            ("footer:even", section.even_page_footer),
        )
        for kind, part in parts:
            if part is None or part.is_linked_to_previous:
                continue
            yield f"{kind}:{index}", part._element, part


def _field_keyword(instruction: str) -> str:
    stripped = instruction.strip()
    if not stripped:
        return ""
    return stripped.split()[0].upper()


def scan_fields(root, location: str) -> list[FieldHit]:
    """按文档顺序扫描复杂域和简单域，标记哪些落在 TOC 结果区内部。

    嵌套靠一个栈跟踪：TOC 的指令文本在 separate 之前就收完了，所以遇到目录内部的
    PAGEREF 时，栈里那层的 keyword 已经能算出来。
    """
    hits: list[FieldHit] = []
    stack: list[dict] = []

    def inside_toc() -> bool:
        return any(_field_keyword("".join(frame["instr"])) == "TOC" for frame in stack)

    for element in root.iter():
        tag = element.tag
        if tag == FLD_CHAR_TAG:
            char_type = element.get(FLD_CHAR_TYPE_ATTR)
            if char_type == "begin":
                stack.append({"instr": [], "in_result": False})
            elif char_type == "separate":
                if stack:
                    stack[-1]["in_result"] = True
            elif char_type == "end":
                if not stack:
                    continue
                frame = stack.pop()
                instruction = "".join(frame["instr"])
                hits.append(
                    FieldHit(
                        keyword=_field_keyword(instruction),
                        instruction=instruction.strip(),
                        location=location,
                        inside_toc=inside_toc(),
                    )
                )
        elif tag == INSTR_TEXT_TAG:
            if stack and not stack[-1]["in_result"]:
                stack[-1]["instr"].append(element.text or "")
        elif tag == FLD_SIMPLE_TAG:
            instruction = element.get(FLD_SIMPLE_INSTR_ATTR) or ""
            hits.append(
                FieldHit(
                    keyword=_field_keyword(instruction),
                    instruction=instruction.strip(),
                    location=location,
                    inside_toc=inside_toc(),
                )
            )
    return hits


def scan_table_samples(document) -> list[TableSampleHit]:
    """扫正文顶层的表头样表标记，以及紧跟其后的那张表。

    只认正文顶层：样表是给模板作者调版式用的，藏进表格单元格或页眉没有意义，
    扫描也会复杂得多。标记后面没有表，或表是空的，都如实记下来交给校验器判。
    """
    samples: list[TableSampleHit] = []
    body = document.element.body
    children = list(body.iterchildren())
    # 段落序号按 iter_paragraphs 的口径数，报错位置才和别的检查对得上。
    index_of: dict[int, int] = {}
    for index, paragraph in enumerate(iter_paragraphs(body, document)):
        index_of[id(paragraph._p)] = index

    for position, child in enumerate(children):
        if not isinstance(child, CT_P):
            continue
        paragraph = Paragraph(child, document)
        text = "".join(run.text for run in visible_runs(paragraph))
        match = TABLE_MARKER_PATTERN.search(text)
        if match is None:
            continue

        following = children[position + 1] if position + 1 < len(children) else None
        table = Table(following, document) if isinstance(following, CT_Tbl) else None
        headers: list[str] = []
        widths: list[int] = []
        if table is not None and table.rows:
            headers = [cell.text.strip() for cell in table.rows[0].cells]
            widths = [
                int(column.width) if column.width is not None else 0
                for column in table.columns
            ]
        samples.append(
            TableSampleHit(
                block=match.group("name").strip(),
                raw=match.group(0),
                paragraph_index=index_of.get(id(child), -1),
                starts_paragraph=text.strip().startswith(match.group(0)),
                has_table=table is not None,
                headers=headers,
                widths=widths,
            )
        )
    return samples


def scan_document(document) -> TemplateScan:
    placeholders: list[PlaceholderHit] = []
    anchors: list[AnchorHit] = []
    fields: list[FieldHit] = []
    unbalanced: list[str] = []

    for location, element, parent in _iter_containers(document):
        fields.extend(scan_fields(element, location))
        for paragraph_index, paragraph in enumerate(iter_paragraphs(element, parent)):
            runs = visible_runs(paragraph)
            text, starts = run_text_map(runs)
            if not text:
                continue

            consumed: list[tuple[int, int]] = []
            for match in PLACEHOLDER_PATTERN.finditer(text):
                name = match.group("name")
                placeholders.append(
                    PlaceholderHit(
                        name=name,
                        raw=match.group(0),
                        location=location,
                        paragraph_index=paragraph_index,
                        span=span_for(starts, match.start(), match.end()),
                        is_well_formed=bool(VALID_PLACEHOLDER_NAME.match(name)),
                    )
                )
                consumed.append((match.start(), match.end()))

            for match in ANCHOR_PATTERN.finditer(text):
                raw = match.group(0)
                name = match.group("name")
                block, _, part = name.partition(":")
                anchors.append(
                    AnchorHit(
                        name=name,
                        block=block.strip(),
                        part=part.strip() or None,
                        raw=raw,
                        location=location,
                        paragraph_index=paragraph_index,
                        owns_paragraph=text.strip() == raw,
                        paragraph_text=text,
                    )
                )
                consumed.append((match.start(), match.end()))

            if _has_unbalanced_braces(text, consumed):
                unbalanced.append(f"{location}#{paragraph_index}")

    return TemplateScan(
        placeholders=placeholders,
        anchors=anchors,
        fields=fields,
        unbalanced_braces=unbalanced,
        table_samples=scan_table_samples(document),
    )


def _has_unbalanced_braces(text: str, consumed: list[tuple[int, int]]) -> bool:
    """匹配对之外还剩 `{{`、`}}` 或残缺锚点前缀时为真。

    残缺的 `{{report_no` 不会被 PLACEHOLDER_PATTERN 命中，若不单独报出来，它会一路
    活到交付文件里。最终 DOCX 校验（设计 §20）也查这个，但模板上传时就该拦住。
    """
    remainder = list(text)
    for start, end in consumed:
        for index in range(start, end):
            remainder[index] = "\0"
    rest = "".join(remainder)
    return "{{" in rest or "}}" in rest or ANCHOR_PREFIX in rest


def scan_template(path: Path) -> TemplateScan:
    return scan_document(Document(str(path)))
