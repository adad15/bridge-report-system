from __future__ import annotations

import base64
from pathlib import Path

from docx import Document
from docx.shared import Inches


PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/p9sAAAAASUVORK5CYII="
)


def write_png(path: Path) -> None:
    path.write_bytes(PNG_1X1)


def add_defect_table(document: Document) -> None:
    document.add_heading("桥梁外观检查", level=1)
    document.add_paragraph("表2.1-1  上部结构病害检查表")
    table = document.add_table(rows=2, cols=6)
    headers = ["构件", "位置", "病害", "数量", "尺寸", "照片编号"]
    values = ["主梁", "第二跨左幅梁底", "裂缝", "1处", "L=0.8m，W=0.12mm", "2.1-1"]
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
        table.rows[1].cells[index].text = values[index]


def add_liaoning_defect_table(document: Document) -> None:
    """真实辽宁国省干线表头形态：含 标度/病害扣分/构件评分 列的两行构件组。"""
    document.add_heading("桥梁外观检查", level=1)
    document.add_paragraph("表2.1-1  上部结构病害检查表")
    headers = [
        "部件名称", "构件编号", "病害位置", "病害类型", "数量",
        "病害特征", "标度", "病害扣分", "构件评分", "照片编号",
    ]
    rows = [
        ["上部承重构件", "2-1#板", "小桩号侧", "蜂窝、麻面", "1处", "S=0.6×0.1m²", "2", "35", "55.81", "2.1-1"],
        ["上部承重构件", "2-1#板", "左侧端部", "剥落、掉角", "1处", "长度：0.5m", "2", "20", "55.81", "2.1-2"],
    ]
    table = document.add_table(rows=1 + len(rows), cols=len(headers))
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for row_index, row in enumerate(rows, start=1):
        for cell_index, value in enumerate(row):
            table.rows[row_index].cells[cell_index].text = value


def add_photo(document: Document, image_path: Path, caption: str = "照片2.1-1 主梁梁底裂缝") -> None:
    document.add_picture(str(image_path), width=Inches(1))
    document.add_paragraph(caption)


def add_weight_table(document: Document) -> None:
    document.add_paragraph("表4.1-1 桥梁部件权重计算表")
    table = document.add_table(rows=2, cols=3)
    rows = [
        ["结构部位", "评价部件", "权重"],
        ["上部结构", "上部承重构件", "0.70"],
    ]
    for row_index, row in enumerate(rows):
        for cell_index, value in enumerate(row):
            table.rows[row_index].cells[cell_index].text = value


def add_rating_table(document: Document) -> None:
    document.add_heading("全桥技术状况综合评定", level=1)
    document.add_paragraph("表4.1-2  总体技术状况评定表")
    table = document.add_table(rows=6, cols=8)
    headers = ["层级", "结构部位", "类别编号", "评价部件", "评分", "权重", "等级", "构件评分"]
    rows = [
        ["全桥", "全桥", "", "全桥", "85.61", "", "2类", ""],
        ["结构分部", "上部结构", "", "上部结构", "87.45", "0.4", "2", ""],
        ["结构分部", "下部结构", "", "下部结构", "86.61", "0.4", "2", ""],
        ["结构分部", "桥面系", "", "桥面系", "79.93", "0.2", "3", ""],
        ["评价部件", "上部结构", "1", "上部承重构件", "86.62", "", "", "3:86.62"],
    ]
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
    for row_index, row in enumerate(rows, start=1):
        for cell_index, value in enumerate(row):
            table.rows[row_index].cells[cell_index].text = value


def create_sample_docx(path: Path, image_path: Path | None = None) -> Path:
    document = Document()
    document.add_paragraph("绕阳河二号桥 定期检测报告")
    add_defect_table(document)
    if image_path is not None:
        add_photo(document, image_path)
    add_rating_table(document)
    document.save(path)
    return path


def create_docx_without_rating_table(path: Path, image_path: Path | None = None) -> Path:
    document = Document()
    document.add_paragraph("绕阳河二号桥 定期检测报告")
    add_defect_table(document)
    if image_path is not None:
        add_photo(document, image_path)
    document.save(path)
    return path
