import asyncio
from pathlib import Path
import zipfile

import pytest
from docx import Document
from docx.shared import Inches
from httpx import ASGITransport, AsyncClient
from pydantic import ValidationError

from bridge_report_tools.contracts.annual_inspection import DataRole, FileRole, SourceType
from bridge_report_tools.importers.defect_tables import parse_defect_tables
from bridge_report_tools.importers.docx_reader import DocxBlocks, DocxTable
from bridge_report_tools.importers.docx_reader import read_docx_blocks
from bridge_report_tools.importers.photo_extractor import extract_and_match_photos, find_photo_captions, media_members
from bridge_report_tools.importers.word_context import ImportMode, WordImportRequest
from bridge_report_tools.importers.word_errors import WordImportError
from bridge_report_tools.importers.word_rules import select_rule_set
from bridge_report_tools.importers.word_importer import parse_word_import
from bridge_report_tools.main import app
from tests.importers.docx_fixtures import (
    add_defect_table,
    add_liaoning_defect_table,
    add_photo,
    add_rating_table,
    create_docx_without_rating_table,
    create_sample_docx,
    write_png,
)


def valid_request(tmp_path: Path) -> WordImportRequest:
    docx_path = tmp_path / "sample.docx"
    docx_path.write_bytes(b"not-a-real-docx-yet")
    photo_dir = tmp_path / "photos"
    photo_dir.mkdir()
    return WordImportRequest(
        docx_path=docx_path,
        temporary_photo_output_dir=photo_dir,
        rule_profile="辽宁国省干线",
        import_mode="已有桥年度导入",
        source_type="软件导出Word",
        file_role="当前年度检测资料",
        data_role="当前年度",
        selected_bridge_system_number="QL-000001",
        selected_bridge_name="绕阳河二号桥",
        inspection_year=2026,
        inspection_date="2026-05-18",
        report_number="Q202605001-JZ-024",
        project_name="绕阳河二号桥2026年度定期检测",
        archived_file_system_number="GDWJ-000001",
        import_record_system_number="DRJL-000001",
    )


def test_word_import_request_accepts_module04_current_year_context(tmp_path: Path) -> None:
    request = valid_request(tmp_path)

    assert request.rule_profile == "辽宁国省干线"
    assert request.import_mode == "已有桥年度导入"
    assert request.source_type == "软件导出Word"
    assert request.file_role == "当前年度检测资料"
    assert request.data_role == "当前年度"
    assert request.docx_path.suffix == ".docx"


def test_word_import_request_accepts_new_bridge_baseline_context(tmp_path: Path) -> None:
    request = valid_request(tmp_path).model_copy(
        update={
            "import_mode": "新桥初始化",
            "source_type": "正式Word",
            "file_role": "历史基线资料",
            "data_role": "历史基线",
            "inspection_year": 2025,
        }
    )

    assert request.import_mode == "新桥初始化"
    assert request.source_type == "正式Word"
    assert request.file_role == "历史基线资料"
    assert request.data_role == "历史基线"
    assert request.inspection_year == 2025


def test_word_import_request_rejects_non_docx(tmp_path: Path) -> None:
    doc_path = tmp_path / "sample.doc"
    doc_path.write_text("old word format", encoding="utf-8")
    photo_dir = tmp_path / "photos"
    photo_dir.mkdir()

    with pytest.raises(ValidationError) as exc_info:
        WordImportRequest(
            docx_path=doc_path,
            temporary_photo_output_dir=photo_dir,
            import_mode="已有桥年度导入",
            source_type="软件导出Word",
            file_role="当前年度检测资料",
            data_role="当前年度",
            selected_bridge_system_number="QL-000001",
            selected_bridge_name="绕阳河二号桥",
            inspection_year=2026,
            inspection_date="2026-05-18",
            report_number="Q202605001-JZ-024",
            project_name="绕阳河二号桥2026年度定期检测",
            archived_file_system_number="GDWJ-000001",
            import_record_system_number="DRJL-000001",
        )

    assert "docx_path" in str(exc_info.value)


def test_word_import_request_rejects_revision_role(tmp_path: Path) -> None:
    request = valid_request(tmp_path)
    with pytest.raises(ValidationError) as exc_info:
        WordImportRequest(
            docx_path=request.docx_path,
            temporary_photo_output_dir=request.temporary_photo_output_dir,
            import_mode=request.import_mode,
            source_type=request.source_type,
            file_role=request.file_role,
            data_role="修订版",
            selected_bridge_system_number=request.selected_bridge_system_number,
            selected_bridge_name=request.selected_bridge_name,
            inspection_year=request.inspection_year,
            inspection_date=request.inspection_date,
            report_number=request.report_number,
            project_name=request.project_name,
            archived_file_system_number=request.archived_file_system_number,
            import_record_system_number=request.import_record_system_number,
        )

    assert "data_role" in str(exc_info.value)


def test_type_aliases_match_contract_literals() -> None:
    source_type: SourceType = "软件导出Word"
    file_role: FileRole = "当前年度检测资料"
    data_role: DataRole = "当前年度"
    import_mode: ImportMode = "已有桥年度导入"

    assert source_type == "软件导出Word"
    assert file_role == "当前年度检测资料"
    assert data_role == "当前年度"
    assert import_mode == "已有桥年度导入"


def test_dynamic_docx_fixture_contains_expected_tables(tmp_path: Path) -> None:
    image_path = tmp_path / "photo.png"
    write_png(image_path)
    docx_path = create_sample_docx(tmp_path / "sample.docx", image_path)

    document = Document(str(docx_path))

    assert len(document.tables) == 2
    assert document.tables[0].rows[0].cells[0].text == "构件"
    assert document.tables[1].rows[0].cells[0].text == "层级"


def test_read_docx_blocks_keeps_table_titles(tmp_path: Path) -> None:
    image_path = tmp_path / "photo.png"
    write_png(image_path)
    docx_path = create_sample_docx(tmp_path / "sample.docx", image_path)

    document = read_docx_blocks(docx_path)

    assert document.paragraph_texts[0] == "绕阳河二号桥 定期检测报告"
    assert document.tables[0].title == "表2.1-1 上部结构病害检查表"
    assert document.tables[0].rows[0] == ["构件", "位置", "病害", "数量", "尺寸", "照片编号"]
    assert document.tables[1].title == "表4.1-2 总体技术状况评定表"


def test_parse_defect_tables_extracts_defect_candidate(tmp_path: Path) -> None:
    image_path = tmp_path / "photo.png"
    write_png(image_path)
    docx_path = create_sample_docx(tmp_path / "sample.docx", image_path)
    document = read_docx_blocks(docx_path)

    rule_set = select_rule_set("辽宁国省干线")
    defects, warnings, errors = parse_defect_tables(document.tables, rule_set)

    assert {warning.code for warning in warnings} == {"liaoning_trunk_defect_table_missing"}
    assert errors == []
    assert len(defects) == 1
    defect = defects[0]
    assert defect.candidate_id == "defect_0001"
    assert defect.source_structure_part == "上部结构"
    assert defect.component_name == "主梁"
    assert defect.defect_location == "第二跨左幅梁底"
    assert defect.defect_type == "裂缝"
    assert defect.quantity_text == "1处"
    assert defect.measurement_text == "L=0.8m，W=0.12mm"
    assert [item.dimension_type for item in defect.measurements] == ["长度", "宽度"]
    assert [item.photo_number for item in defect.photo_references] == ["2.1-1"]
    assert all(item.resolution == "pending" for item in defect.photo_references)
    assert defect.source_ref.table_title == "表2.1-1 上部结构病害检查表"


def test_parse_defect_tables_ignores_non_defect_table_with_disease_title() -> None:
    table = DocxTable(
        index=0,
        title="主要病害及技术状况评定表",
        chapter="第四章 全桥技术状况综合评定",
        rows=[
            ["层级", "结构部位", "类别编号", "评价部件", "评分", "权重", "等级", "构件评分"],
            ["全桥", "全桥", "", "全桥", "85.61", "", "2类", ""],
        ],
    )

    rule_set = select_rule_set("辽宁国省干线")
    defects, warnings, errors = parse_defect_tables([table], rule_set)

    assert defects == []
    assert warnings == []
    assert len(errors) == 1
    assert errors[0].code == "defect_tables_not_found"


def test_parse_defect_tables_uses_liaoning_table_numbers() -> None:
    rule_set = select_rule_set("辽宁国省干线")
    table = DocxTable(
        index=0,
        title="表2.2-1  下部结构病害检查表",
        chapter=None,
        rows=[
            ["构件", "位置", "病害", "数量", "尺寸", "照片编号"],
            ["桥台", "0#台左侧翼墙", "勾缝砂浆脱落", "1处", "L=1m", "2.2-1"],
        ],
    )

    defects, warnings, errors = parse_defect_tables([table], rule_set)

    assert warnings
    assert {warning.code for warning in warnings} == {"liaoning_trunk_defect_table_missing"}
    assert errors == []
    assert len(defects) == 1
    assert defects[0].source_structure_part == "下部结构"
    assert [item.photo_number for item in defects[0].photo_references] == ["2.2-1"]


def test_parse_defect_tables_maps_real_liaoning_header_aliases() -> None:
    rule_set = select_rule_set("辽宁国省干线")
    table = DocxTable(
        index=0,
        title="表2.1-1 上部结构病害检查表",
        chapter=None,
        rows=[
            [
                "序号",
                "部件 名称",
                "构件 编号",
                "病害位置",
                "病害 类型",
                "病害特征",
                "标度",
                "病害扣分",
                "构件评分",
                "照片编号",
            ],
            [
                "1",
                "上部承重构件",
                "1-1#板",
                "梁底",
                "横向裂缝",
                "L=6m，W=0.2mm",
                "2",
                "35",
                "55.81",
                "照片2.1-1",
            ],
        ],
    )

    defects, warnings, errors = parse_defect_tables([table], rule_set)

    assert errors == []
    assert {warning.code for warning in warnings} == {"liaoning_trunk_defect_table_missing"}
    assert len(defects) == 1
    defect = defects[0]
    assert defect.component_name == "上部承重构件"
    assert defect.component_number == "1-1#板"
    assert defect.defect_location == "梁底"
    assert defect.defect_type == "横向裂缝"
    assert defect.measurement_text == "L=6m，W=0.2mm"
    assert [item.dimension_type for item in defect.measurements] == ["长度", "宽度"]
    assert [item.photo_number for item in defect.photo_references] == ["2.1-1"]
    assert defect.severity is None


def test_parse_defect_tables_reports_missing_defect_table(tmp_path: Path) -> None:
    image_path = tmp_path / "photo.png"
    write_png(image_path)
    docx_path = create_sample_docx(tmp_path / "sample.docx", image_path)
    document = read_docx_blocks(docx_path)

    rule_set = select_rule_set("辽宁国省干线")
    defects, warnings, errors = parse_defect_tables(document.tables[1:], rule_set)

    assert defects == []
    assert warnings == []
    assert len(errors) == 1
    assert errors[0].code == "defect_tables_not_found"


def test_parse_word_import_outputs_contract_data_and_photo_files(tmp_path: Path) -> None:
    image_path = tmp_path / "photo.png"
    write_png(image_path)
    request = valid_request(tmp_path)
    docx_path = create_sample_docx(request.docx_path, image_path)
    request = request.model_copy(
        update={
            "docx_path": docx_path,
            "temporary_photo_output_dir": tmp_path / "out",
        }
    )

    assert request.rule_profile == "辽宁国省干线"

    response = parse_word_import(request)

    assert response.temporary_photo_files == ["photo_0001.png"]
    data = response.data
    assert data.contract.name == "BridgeAnnualInspectionData"
    assert data.contract.version == "5.0"
    assert data.contract.parser_name == "word_importer"
    assert data.import_context.source_type == "软件导出Word"
    assert data.import_context.file_role == "当前年度检测资料"
    assert data.bridge_check.selected_bridge_system_number == "QL-000001"
    assert data.bridge_check.extracted_bridge_name == "绕阳河二号桥"
    assert data.bridge_check.match_status == "匹配"
    assert data.inspection.inspection_year == 2026
    assert data.inspection.report_number == "Q202605001-JZ-024"
    assert len(data.defects) == 1
    assert data.defects[0].candidate_id == "defect_0001"
    assert all(item.group_review_status == "待确认" for item in data.defects)
    assert all(
        all(reference.resolution == "pending" for reference in item.photo_references)
        for item in data.defects
    )
    assert len(data.photos) == 1
    assert data.photos[0].linked_defect_candidate_id == "defect_0001"
    assert data.contract.version == "5.0"
    assert not hasattr(data, "ratings")
    assert data.comparison_candidates == []
    assert data.report_text_candidates == []
    assert data.errors == []


def test_parse_word_import_keeps_defect_table_missing_as_contract_error(tmp_path: Path) -> None:
    request = valid_request(tmp_path)
    document_obj = Document()
    document_obj.add_paragraph("绕阳河二号桥 定期检测报告")
    add_rating_table(document_obj)
    docx_path = request.docx_path
    document_obj.save(docx_path)
    request = request.model_copy(
        update={
            "docx_path": docx_path,
            "temporary_photo_output_dir": tmp_path / "out",
        }
    )

    response = parse_word_import(request)

    assert response.data.defects == []
    assert response.data.photos == []
    assert response.data.errors[0].code == "defect_tables_not_found"
    assert response.data.contract.version == "5.0"
    assert not hasattr(response.data, "ratings")


def test_parse_word_import_succeeds_when_rating_table_missing(tmp_path: Path) -> None:
    image_path = tmp_path / "photo.png"
    write_png(image_path)
    request = valid_request(tmp_path)
    docx_path = create_docx_without_rating_table(request.docx_path, image_path)
    request = request.model_copy(
        update={
            "docx_path": docx_path,
            "temporary_photo_output_dir": tmp_path / "out",
        }
    )

    response = parse_word_import(request)

    assert response.data.contract.version == "5.0"
    assert len(response.data.defects) == 1
    assert not hasattr(response.data, "ratings")
    assert "rating_table_not_found" not in {
        warning.code for warning in response.data.warnings + response.data.errors
    }


def test_parse_word_endpoint_returns_contract_data(tmp_path: Path) -> None:
    image_path = tmp_path / "photo.png"
    write_png(image_path)
    request = valid_request(tmp_path)
    docx_path = create_sample_docx(request.docx_path, image_path)
    request = request.model_copy(
        update={
            "docx_path": docx_path,
            "temporary_photo_output_dir": tmp_path / "out",
        }
    )

    async def post_parse():
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://testserver") as client:
            return await client.post(
                "/imports/word/parse",
                json=request.model_dump(mode="json"),
            )

    response = asyncio.run(post_parse())

    assert response.status_code == 200
    payload = response.json()
    assert payload["data"]["contract"]["name"] == "BridgeAnnualInspectionData"
    assert payload["data"]["contract"]["parser_name"] == "word_importer"
    assert payload["data"]["defects"][0]["candidate_id"] == "defect_0001"
    assert payload["temporary_photo_files"] == ["photo_0001.png"]


def test_parse_word_endpoint_accepts_document_without_rating_table(tmp_path: Path) -> None:
    image_path = tmp_path / "photo.png"
    write_png(image_path)
    request = valid_request(tmp_path)
    docx_path = create_docx_without_rating_table(request.docx_path, image_path)
    request = request.model_copy(
        update={
            "docx_path": docx_path,
            "temporary_photo_output_dir": tmp_path / "out",
        }
    )

    async def post_parse():
        transport = ASGITransport(app=app)
        async with AsyncClient(transport=transport, base_url="http://testserver") as client:
            return await client.post(
                "/imports/word/parse",
                json=request.model_dump(mode="json"),
            )

    response = asyncio.run(post_parse())

    assert response.status_code == 200
    payload = response.json()
    assert payload["data"]["contract"]["version"] == "5.0"
    assert "ratings" not in payload["data"]


def test_parse_defect_tables_keeps_row_level_warnings() -> None:
    rule_set = select_rule_set("辽宁国省干线")
    table = DocxTable(
        index=0,
        title="表2.1-1 上部结构病害检查表",
        chapter=None,
        rows=[
            ["构件", "位置", "病害", "数量", "尺寸", "照片编号"],
            ["主梁", "第二跨左幅梁底", "破损", "1处", "局部破损，约20左右", ""],
        ],
    )

    defects, warnings, errors = parse_defect_tables([table], rule_set)

    assert {warning.code for warning in warnings} == {"liaoning_trunk_defect_table_missing"}
    assert errors == []
    assert len(defects) == 1
    warning_codes = [warning.code for warning in defects[0].warnings]
    assert "measurement_parse_low_confidence" in warning_codes
    assert "photo_number_missing" not in warning_codes


def test_extract_and_match_photos_links_caption_to_defect(tmp_path: Path) -> None:
    image_path = tmp_path / "photo.png"
    write_png(image_path)
    docx_path = create_sample_docx(tmp_path / "sample.docx", image_path)
    document = read_docx_blocks(docx_path)
    defects, _, _ = parse_defect_tables(document.tables, select_rule_set("辽宁国省干线"))
    photo_output_dir = tmp_path / "out"

    photos, temporary_files, warnings = extract_and_match_photos(
        docx_path, document, defects, photo_output_dir, select_rule_set("辽宁国省干线")
    )

    assert warnings == []
    assert temporary_files == ["photo_0001.png"]
    assert (photo_output_dir / "photo_0001.png").exists()
    assert len(photos) == 1
    assert photos[0].photo_number == "2.1-1"
    assert photos[0].linked_defect_candidate_id == "defect_0001"
    assert not hasattr(photos[0], "match_status")
    assert not hasattr(photos[0], "review_status")
    assert photos[0].extracted_file.temporary_file_name == "photo_0001.png"
    assert photos[0].extracted_file.original_caption == "照片2.1-1 主梁梁底裂缝"


def test_find_photo_captions_ignores_dates_without_caption_prefix() -> None:
    rule_set = select_rule_set("辽宁国省干线")
    document = DocxBlocks(
        paragraph_texts=[
            "检测日期 2026-05-18",
            "照片 2026-05-18",
            "照片1-1 桥梁正面照",
            "照片2.1-1 主梁裂缝",
        ],
        tables=[],
    )

    captions = find_photo_captions(document, rule_set)

    assert [(caption.number, caption.is_defect_photo) for caption in captions] == [
        ("1-1", False),
        ("2.1-1", True),
    ]


def test_liaoning_photo_caption_normalizes_compact_table_text() -> None:
    rule_set = select_rule_set("辽宁国省干线")

    caption = rule_set.parse_photo_caption("照片2.21 0#台左侧翼墙勾缝砂浆脱落")
    two_digit_caption = rule_set.parse_photo_caption("照片2.110 2-1#铰缝勾缝砂浆脱落")

    assert caption is not None
    assert caption.number == "2.2-1"
    assert caption.is_defect_photo is True
    assert two_digit_caption is not None
    assert two_digit_caption.number == "2.1-10"
    assert two_digit_caption.is_defect_photo is True


def test_extract_and_match_photos_links_table_cell_caption_below_image(tmp_path: Path) -> None:
    rule_set = select_rule_set("辽宁国省干线")
    image_path = tmp_path / "photo.png"
    write_png(image_path)
    docx_path = tmp_path / "sample.docx"
    document_obj = Document()
    add_defect_table(document_obj)
    gallery = document_obj.add_table(rows=1, cols=1)
    cell = gallery.rows[0].cells[0]
    cell.paragraphs[0].add_run().add_picture(str(image_path), width=Inches(1))
    cell.add_paragraph("照片2.1-1 主梁梁底裂缝")
    add_rating_table(document_obj)
    document_obj.save(docx_path)
    document = read_docx_blocks(docx_path)
    defects, _, _ = parse_defect_tables(document.tables, rule_set)

    photos, temporary_files, warnings = extract_and_match_photos(
        docx_path,
        document,
        defects,
        tmp_path / "out",
        rule_set,
    )

    assert warnings == []
    assert temporary_files == ["photo_0001.png"]
    assert len(photos) == 1
    assert photos[0].photo_number == "2.1-1"
    assert photos[0].linked_defect_candidate_id == "defect_0001"
    assert not hasattr(photos[0], "match_status")
    assert photos[0].extracted_file.original_caption == "照片2.1-1 主梁梁底裂缝"


def test_extract_and_match_photos_skips_overview_photos_in_candidates(tmp_path: Path) -> None:
    rule_set = select_rule_set("辽宁国省干线")
    image_path = tmp_path / "photo.png"
    write_png(image_path)
    docx_path = tmp_path / "sample.docx"
    document_obj = Document()
    add_defect_table(document_obj)
    add_photo(document_obj, image_path, "照片1-1 桥梁正面照")
    add_photo(document_obj, image_path, "照片2.1-1 主梁梁底裂缝")
    add_rating_table(document_obj)
    document_obj.save(docx_path)
    document = read_docx_blocks(docx_path)
    defects, _, _ = parse_defect_tables(document.tables, select_rule_set("辽宁国省干线"))

    photos, temporary_files, warnings = extract_and_match_photos(
        docx_path,
        document,
        defects,
        tmp_path / "out",
        rule_set,
    )

    assert warnings == []
    assert temporary_files == ["photo_0001.png", "photo_0002.png"]
    assert len(photos) == 1
    assert photos[0].photo_number == "2.1-1"
    assert photos[0].extracted_file.temporary_file_name == "photo_0002.png"


def test_extract_and_match_photos_keeps_unreferenced_photo_warning(tmp_path: Path) -> None:
    image_path = tmp_path / "photo.png"
    write_png(image_path)
    docx_path = tmp_path / "sample.docx"
    document_obj = Document()
    add_defect_table(document_obj)
    add_photo(document_obj, image_path, "照片2.1-3 桥面铺装局部破损")
    add_rating_table(document_obj)
    document_obj.save(docx_path)
    document = read_docx_blocks(docx_path)
    defects, _, _ = parse_defect_tables(document.tables, select_rule_set("辽宁国省干线"))

    photos, temporary_files, warnings = extract_and_match_photos(
        docx_path, document, defects, tmp_path / "out", select_rule_set("辽宁国省干线")
    )

    assert warnings == []
    assert temporary_files == ["photo_0001.png"]
    assert photos[0].photo_number == "2.1-3"
    assert photos[0].linked_defect_candidate_id is None
    assert not hasattr(photos[0], "match_status")
    assert photos[0].warnings[0].code == "photo_not_referenced_by_defect"


def test_media_members_uses_document_body_relationship_order(tmp_path: Path) -> None:
    docx_path = tmp_path / "ordered.docx"
    with zipfile.ZipFile(docx_path, "w") as archive:
        archive.writestr(
            "word/document.xml",
            """
            <w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
                xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
                xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
                <w:body>
                    <w:p><w:r><w:drawing><a:blip r:embed="rId2" /></w:drawing></w:r></w:p>
                    <w:p><w:r><w:drawing><a:blip r:embed="rId1" /></w:drawing></w:r></w:p>
                </w:body>
            </w:document>
            """,
        )
        archive.writestr(
            "word/_rels/document.xml.rels",
            """
            <Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
                <Relationship Id="rId1"
                    Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
                    Target="media/image10.png" />
                <Relationship Id="rId2"
                    Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"
                    Target="media/image2.png" />
            </Relationships>
            """,
        )
        archive.writestr("word/media/image10.png", b"image10")
        archive.writestr("word/media/image2.png", b"image2")
        archive.writestr("word/media/logo.png", b"logo")

    assert media_members(docx_path) == ["word/media/image2.png", "word/media/image10.png"]


def test_read_docx_blocks_keeps_nested_tables_as_tables(tmp_path: Path) -> None:
    docx_path = tmp_path / "nested.docx"
    document_obj = Document()
    outer = document_obj.add_table(rows=1, cols=1)
    cell = outer.rows[0].cells[0]
    cell.add_paragraph("表2.1-1  上部结构病害检查表")
    nested = cell.add_table(rows=2, cols=6)
    headers = ["构件", "位置", "病害", "数量", "尺寸", "照片编号"]
    values = ["主梁", "第二跨左幅梁底", "裂缝", "1处", "L=0.8m", "2.1-1"]
    for index, header in enumerate(headers):
        nested.rows[0].cells[index].text = header
        nested.rows[1].cells[index].text = values[index]
    document_obj.save(docx_path)

    document = read_docx_blocks(docx_path)

    assert any(table.title == "表2.1-1 上部结构病害检查表" for table in document.tables)
    assert any(table.rows[0] == headers for table in document.tables)


def test_read_docx_blocks_reads_tables_inside_custom_xml_blocks(tmp_path: Path) -> None:
    from docx.oxml import OxmlElement

    docx_path = tmp_path / "customxml.docx"
    document_obj = Document()
    caption = document_obj.add_paragraph("表2.1-1  上部结构病害检查表")
    table = document_obj.add_table(rows=2, cols=6)
    headers = ["构件", "位置", "病害", "数量", "尺寸", "照片编号"]
    values = ["主梁", "第二跨左幅梁底", "裂缝", "1处", "L=0.8m", "2.1-1"]
    for index, header in enumerate(headers):
        table.rows[0].cells[index].text = header
        table.rows[1].cells[index].text = values[index]
    wrapper = OxmlElement("w:customXml")
    body = document_obj.element.body
    body.insert(list(body).index(caption._p), wrapper)
    wrapper.append(caption._p)
    wrapper.append(table._tbl)
    document_obj.save(docx_path)

    document = read_docx_blocks(docx_path)

    assert "表2.1-1 上部结构病害检查表" in document.paragraph_texts
    assert any(table.title == "表2.1-1 上部结构病害检查表" for table in document.tables)
    assert any(table.rows[0] == headers for table in document.tables)


def test_docx_diagnostics_finds_text_in_document_xml(tmp_path: Path) -> None:
    from bridge_report_tools.importers.docx_diagnostics import find_text_locations

    docx_path = tmp_path / "diagnostic.docx"
    document_obj = Document()
    document_obj.add_paragraph("表4.1-2  总体技术状况评定表")
    document_obj.save(docx_path)

    locations = find_text_locations(docx_path, ["表4.1-2"])

    assert locations["表4.1-2"] == ["word/document.xml"]


def test_extract_and_match_photos_unmatched_defect_warning_is_idempotent(tmp_path: Path) -> None:
    image_path = tmp_path / "photo.png"
    write_png(image_path)
    docx_path = tmp_path / "sample.docx"
    document_obj = Document()
    add_defect_table(document_obj)
    add_photo(document_obj, image_path, "照片2.1-3 桥面铺装局部破损")
    add_rating_table(document_obj)
    document_obj.save(docx_path)
    document = read_docx_blocks(docx_path)
    defects, _, _ = parse_defect_tables(document.tables, select_rule_set("辽宁国省干线"))

    rule_set = select_rule_set("辽宁国省干线")
    extract_and_match_photos(docx_path, document, defects, tmp_path / "out", rule_set)
    extract_and_match_photos(docx_path, document, defects, tmp_path / "out", rule_set)

    warning_codes = [warning.code for warning in defects[0].warnings]
    assert warning_codes.count("photo_number_unmatched") == 1


def test_parse_word_import_ignores_rating_and_deduction_columns(tmp_path: Path) -> None:
    request = valid_request(tmp_path)
    document_obj = Document()
    document_obj.add_paragraph("绕阳河二号桥 定期检测报告")
    add_liaoning_defect_table(document_obj)
    add_rating_table(document_obj)
    docx_path = tmp_path / "liaoning.docx"
    document_obj.save(docx_path)
    request = request.model_copy(update={"docx_path": docx_path})

    response = parse_word_import(request)

    dumped = response.data.model_dump(mode="json")
    assert "ratings" not in dumped
    assert all("defect_deduction" not in defect for defect in dumped["defects"])
    assert [defect.component_number for defect in response.data.defects] == ["2-1#板", "2-1#板"]
    assert [defect.defect_scale for defect in response.data.defects] == [2, 2]
