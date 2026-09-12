from __future__ import annotations

from pathlib import Path

from docx import Document

from bridge_report_tools.reports.docx_scan import (
    LOCATION_BODY,
    scan_document,
    scan_template,
    visible_runs,
)
from tests.reports.template_fixtures import (
    CORE_ANCHORS,
    add_complex_field,
    add_simple_field,
    add_split_text,
    add_toc,
    minimal_valid_template,
)


def test_placeholder_split_across_runs_is_found_with_full_span() -> None:
    document = Document()
    paragraph = document.add_paragraph()
    add_split_text(paragraph, "编号：", "{{repo", "rt", "_no}}", " 结束")

    scan = scan_document(document)

    assert [hit.name for hit in scan.placeholders] == ["report_no"]
    hit = scan.placeholders[0]
    assert hit.raw == "{{report_no}}"
    assert hit.is_well_formed
    # 跨越 run 1..3；替换时要能把整段抹掉而不是只改中间一个 run。
    assert hit.span.start_run == 1
    assert hit.span.start_offset == 0
    assert hit.span.end_run == 3
    assert hit.span.end_offset == len("_no}}")


def test_placeholder_inside_single_run_maps_to_exact_offsets() -> None:
    document = Document()
    document.add_paragraph("前缀 {{bridge_name}} 后缀")

    scan = scan_document(document)

    span = scan.placeholders[0].span
    assert span.start_run == 0
    assert span.end_run == 0
    assert span.start_offset == len("前缀 ")
    assert span.end_offset == len("前缀 {{bridge_name}}")


def test_non_ascii_placeholder_name_is_reported_as_malformed() -> None:
    document = Document()
    document.add_paragraph("{{桥梁名称}}")

    scan = scan_document(document)

    assert len(scan.placeholders) == 1
    assert not scan.placeholders[0].is_well_formed


def test_unclosed_placeholder_is_reported_as_unbalanced() -> None:
    document = Document()
    document.add_paragraph("{{report_no 少了右括号")

    scan = scan_document(document)

    assert scan.placeholders == []
    assert scan.unbalanced_braces == [f"{LOCATION_BODY}#0"]


def test_anchor_owning_its_paragraph_is_marked() -> None:
    document = Document()
    document.add_paragraph("[[REPORT:DEFECT_TABLES]]")
    document.add_paragraph("正文 [[REPORT:DEFECT_PHOTOS]] 混排")

    scan = scan_document(document)

    assert [hit.name for hit in scan.anchors] == ["DEFECT_TABLES", "DEFECT_PHOTOS"]
    assert scan.anchors[0].owns_paragraph
    assert not scan.anchors[1].owns_paragraph


def test_anchor_splits_block_and_structure_part() -> None:
    document = Document()
    document.add_paragraph("[[REPORT:DEFECT_TABLES:SUPERSTRUCTURE]]")
    document.add_paragraph("[[REPORT:CONCLUSION]]")

    scan = scan_document(document)

    assert (scan.anchors[0].block, scan.anchors[0].part) == (
        "DEFECT_TABLES",
        "SUPERSTRUCTURE",
    )
    assert (scan.anchors[1].block, scan.anchors[1].part) == ("CONCLUSION", None)


def test_anchor_with_surrounding_whitespace_still_owns_paragraph() -> None:
    document = Document()
    add_split_text(document.add_paragraph(), "  [[REPORT:", "CONCLUSION", "]]  ")

    scan = scan_document(document)

    assert scan.anchors[0].name == "CONCLUSION"
    assert scan.anchors[0].owns_paragraph


def test_placeholders_in_table_cells_are_scanned_once_despite_merge() -> None:
    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "{{route_code}}"
    table.rows[0].cells[1].text = "{{route_name}}"
    # 合并后 Table.rows[].cells 会把同一个 w:tc 返回两次；扫描不能跟着数两次。
    table.rows[0].cells[0].merge(table.rows[0].cells[1])

    scan = scan_document(document)

    assert [hit.name for hit in scan.placeholders] == ["route_code", "route_name"]


def test_placeholder_in_header_is_scanned_with_header_location() -> None:
    document = Document()
    header = document.sections[0].header
    header.is_linked_to_previous = False
    header.paragraphs[0].text = "{{bridge_name}}"

    scan = scan_document(document)

    assert [hit.location for hit in scan.placeholders] == ["header:1"]


def test_toc_internal_pageref_is_marked_inside_toc() -> None:
    document = Document()
    add_toc(document.add_paragraph())

    scan = scan_document(document)

    by_keyword = {hit.keyword: hit for hit in scan.fields}
    assert set(by_keyword) == {"TOC", "PAGEREF"}
    assert not by_keyword["TOC"].inside_toc
    assert by_keyword["PAGEREF"].inside_toc


def test_standalone_pageref_is_not_marked_inside_toc() -> None:
    document = Document()
    add_complex_field(document.add_paragraph(), " PAGEREF _Ref9 \\h ", "12")

    scan = scan_document(document)

    assert [hit.keyword for hit in scan.fields] == ["PAGEREF"]
    assert not scan.fields[0].inside_toc


def test_simple_field_is_scanned() -> None:
    document = Document()
    add_simple_field(document.add_paragraph(), " PAGE ", "1")

    scan = scan_document(document)

    assert [hit.keyword for hit in scan.fields] == ["PAGE"]


def test_seq_field_keyword_is_reported() -> None:
    document = Document()
    add_complex_field(document.add_paragraph(), " SEQ 表 \\* ARABIC ", "1")

    scan = scan_document(document)

    assert scan.fields[0].keyword == "SEQ"


def test_field_instruction_text_is_not_read_as_placeholder() -> None:
    """域指令写在 w:instrText 里，不是可见文字；不能被当成占位符扫出来。"""
    document = Document()
    add_complex_field(document.add_paragraph(), " REF {{report_no}} \\h ", "x")

    scan = scan_document(document)

    assert scan.placeholders == []
    assert scan.fields[0].keyword == "REF"


def test_visible_runs_includes_hyperlink_runs() -> None:
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    document = Document()
    paragraph = document.add_paragraph()
    paragraph._p.append(
        parse_xml(
            f'<w:hyperlink {nsdecls("w")} w:anchor="_Toc1">'
            f"<w:r><w:t>{{{{bridge_name}}}}</w:t></w:r>"
            f"</w:hyperlink>"
        )
    )

    assert len(visible_runs(paragraph)) == 1
    assert [hit.name for hit in scan_document(document).placeholders] == ["bridge_name"]


def test_scan_template_reads_from_disk(tmp_path: Path) -> None:
    scan = scan_template(minimal_valid_template(tmp_path / "template.docx"))

    assert [hit.name for hit in scan.anchors] == list(CORE_ANCHORS)
    assert {hit.name for hit in scan.placeholders} == {"report_no", "bridge_name"}
    assert scan.unbalanced_braces == []
