"""最终 DOCX 校验（设计 §20）。

守的是同一件事：**一份看着正常、实际残缺的报告绝不能交付。** 下面每一条都是在
Word 里打开不会报错、读者却会拿到错报告的情形。
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from docx import Document

from bridge_report_tools.reports.contract import PERIODIC_INSPECTION_V1, anchor_text
from bridge_report_tools.reports.errors import ReportTemplateError
from bridge_report_tools.reports.output_validator import validate_output

from tests.reports.context_fixtures import defect_row, part, photo, png
from tests.reports.context_fixtures import context as report_context
from tests.reports.template_fixtures import (
    CORE_ANCHORS,
    add_complex_field,
    build_builder_template,
)


ALL_BLOCKS = list(CORE_ANCHORS)


def good_report(tmp_path: Path) -> Path:
    """一份装配完整、带照片的成品。"""
    archive = tmp_path / "archive"
    (archive / "photos").mkdir(parents=True)
    png(archive / "photos/a.png", 800, 600)
    from bridge_report_tools.reports.docx_builder import build_report

    ctx = report_context(
        parts=[
            part(
                "SUPERSTRUCTURE",
                "上部结构",
                rows=[defect_row(1, photo_numbers=["照片2.1-1"])],
                photos=[photo("照片2.1-1", "photos/a.png")],
            )
        ]
    )
    template = build_builder_template(tmp_path / "template.docx")
    result = build_report(template, ctx, tmp_path / "job" / "report.docx", archive)
    return result.output_path


def check(path: Path, job: Path, blocks=None):
    return validate_output(
        path, job, PERIODIC_INSPECTION_V1, ALL_BLOCKS if blocks is None else blocks
    )


def codes(result) -> list[str]:
    return [issue.code for issue in result.issues]


def test_a_complete_report_passes(tmp_path: Path) -> None:
    report = good_report(tmp_path)

    result = check(report, tmp_path / "job")

    assert result.status == "valid", codes(result)
    assert result.section_count >= 1
    assert result.image_count == 1


def test_leftover_anchor_is_rejected(tmp_path: Path) -> None:
    """锚点没被替换掉，在 Word 里就是一行 [[REPORT:...]]，报告却"打得开"。"""
    report = good_report(tmp_path)
    document = Document(str(report))
    document.add_paragraph(anchor_text("CONCLUSION"))
    document.save(str(report))

    result = check(report, tmp_path / "job")

    assert "report_output_anchor_left" in codes(result)


def test_half_eaten_anchor_is_rejected_too(tmp_path: Path) -> None:
    """装配中途出错可能只留下半截标记，扫描器不认它，但同样不能交付。"""
    report = good_report(tmp_path)
    document = Document(str(report))
    document.add_paragraph("[[REPORT:CONCLUSION")
    document.save(str(report))

    result = check(report, tmp_path / "job")

    assert "report_output_anchor_left" in codes(result)


def test_leftover_placeholder_is_rejected(tmp_path: Path) -> None:
    report = good_report(tmp_path)
    document = Document(str(report))
    document.add_paragraph("报告编号：{{report_no}}")
    document.save(str(report))

    result = check(report, tmp_path / "job")

    assert "report_output_placeholder_left" in codes(result)


def test_missing_required_block_is_rejected(tmp_path: Path) -> None:
    """装配层说漏了一块，成品文件本身看不出来，只能靠对账。"""
    report = good_report(tmp_path)

    result = check(report, tmp_path / "job", blocks=[b for b in ALL_BLOCKS if b != "CONCLUSION"])

    assert "report_output_block_missing" in codes(result)


def test_missing_media_file_is_rejected(tmp_path: Path) -> None:
    """图片关系还在、媒体文件没了：Word 打开显示一个红叉，不报错。"""
    report = good_report(tmp_path)
    rebuilt = tmp_path / "job" / "broken.docx"
    with zipfile.ZipFile(report) as source, zipfile.ZipFile(rebuilt, "w") as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename.startswith("word/media/"):
                data = b""
            target.writestr(item, data)

    result = check(rebuilt, tmp_path / "job")

    assert "report_output_media_missing" in codes(result)


def test_disallowed_field_is_rejected(tmp_path: Path) -> None:
    """刷域后冒出来的 SEQ 一律拒绝——表号图号只能由生成器确定性写入（§7.6）。"""
    report = good_report(tmp_path)
    document = Document(str(report))
    add_complex_field(document.add_paragraph(), " SEQ 表 \\* ARABIC ", "1")
    document.save(str(report))

    result = check(report, tmp_path / "job")

    assert "report_output_field_not_allowed" in codes(result)


def test_toc_internal_pageref_is_allowed(tmp_path: Path) -> None:
    """目录刷完后 Word 自己写的 PAGEREF 属于白名单，不能因此判成品不合格。"""
    report = good_report(tmp_path)
    document = Document(str(report))
    add_complex_field(
        document.add_paragraph(),
        ' TOC \\o "1-3" \\h ',
        "",
        nested=[(" PAGEREF _Toc1 \\h ", "3")],
    )
    document.save(str(report))

    result = check(report, tmp_path / "job")

    assert result.status == "valid", codes(result)


def test_output_outside_the_job_directory_is_rejected(tmp_path: Path) -> None:
    """下载按任务 ID 取路径；路径能指到任务目录之外就等于开了读任意文件的口子。"""
    report = good_report(tmp_path)

    result = check(report, tmp_path / "another-job")

    assert "report_output_path_outside_job" in codes(result)


def test_a_file_that_is_not_a_zip_is_rejected(tmp_path: Path) -> None:
    job = tmp_path / "job"
    job.mkdir()
    broken = job / "report.docx"
    broken.write_bytes(b"not a docx at all")

    result = check(broken, job)

    assert codes(result) == ["report_output_not_ooxml"]


def test_a_missing_file_raises_instead_of_returning_issues(tmp_path: Path) -> None:
    """文件根本不在，连明细都无从谈起——这是任务失败，不是"校验不通过"。"""
    with pytest.raises(ReportTemplateError) as raised:
        check(tmp_path / "job" / "nope.docx", tmp_path / "job")

    assert raised.value.code == "report_output_missing"
