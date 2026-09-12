"""Word / WPS 域更新器的真机测试。

默认跳过：这些用例会真的启动 Word 和 WPS。设 BRIDGE_REPORT_OFFICE_TESTS=1 再跑，
对应设计 §27 第 3 步的阶段门。
"""

from __future__ import annotations

import os
import threading
import time
from pathlib import Path

import pytest
from docx import Document
from docx.enum.section import WD_SECTION

from bridge_report_tools.reports.field_updater import (
    UPDATER_MICROSOFT_WORD,
    UPDATER_WPS_WRITER,
    FieldUpdateError,
    microsoft_word_updater,
    probe_all,
    wps_writer_updater,
)
from tests.reports.template_fixtures import add_simple_field, add_toc


pytestmark = pytest.mark.skipif(
    os.environ.get("BRIDGE_REPORT_OFFICE_TESTS") != "1",
    reason="需要 BRIDGE_REPORT_OFFICE_TESTS=1，且本机装有 Word 或 WPS",
)

HEADINGS = ("第一章 桥梁概况", "第二章 病害检查", "第三章 技术状况评定")


def build_multipage_document(path: Path) -> Path:
    """一份有目录、有页码、跨多页的文档，用来验证域是否真的被算出来了。"""
    document = Document()
    document.add_paragraph("目录")
    add_toc(document.add_paragraph())
    document.add_page_break()
    for title in HEADINGS:
        document.add_heading(title, level=1)
        for index in range(30):
            document.add_paragraph(f"正文占位段落 {index}，用于把文档撑到多页。")
        document.add_page_break()

    footer = document.sections[0].footer
    footer.is_linked_to_previous = False
    add_simple_field(footer.paragraphs[0], " PAGE ", "1")
    add_simple_field(footer.paragraphs[0], " NUMPAGES ", "1")

    document.save(str(path))
    return path


def build_sectioned_document(path: Path, sections: int = 7) -> Path:
    """多节、每节都有页眉页脚的文档——真实报告就是这个形状（封面、声明、目录、
    正文、两个附录各一节）。

    专门守住一条：域更新必须走得完。串 NextStoryRange 遍历故事区时，本机的 WPS
    在页眉页脚那条链上不收敛，同一个页眉无限重复，更新永远结束不了；单节文档没有
    这条链，测不出来。
    """
    document = Document()
    for number in range(sections):
        if number:
            document.add_section(WD_SECTION.NEW_PAGE)
        section = document.sections[number]
        section.header.is_linked_to_previous = False
        section.header.paragraphs[0].text = f"第 {number + 1} 节页眉"
        section.footer.is_linked_to_previous = False
        add_simple_field(section.footer.paragraphs[0], " PAGE ", "1")
        add_simple_field(section.footer.paragraphs[0], " NUMPAGES ", "1")
        for index in range(20):
            document.add_paragraph(f"第 {number + 1} 节正文段落 {index}。")
    document.save(str(path))
    return path


def footer_text(path: Path) -> list[str]:
    return [
        "".join(paragraph.text for paragraph in section.footer.paragraphs)
        for section in Document(str(path)).sections
    ]


def document_text(path: Path) -> str:
    document = Document(str(path))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _updater_or_skip(factory):
    updater = factory(timeout_seconds=180)
    probe = updater.probe()
    if not probe.available:
        pytest.skip(f"{updater.name} 不可用：{probe.detail}")
    return updater


def test_probe_reports_at_least_one_updater() -> None:
    results = probe_all(timeout_seconds=120)

    assert {result.updater for result in results} == {
        UPDATER_MICROSOFT_WORD,
        UPDATER_WPS_WRITER,
    }
    assert any(result.available for result in results), [r.detail for r in results]


@pytest.mark.parametrize(
    "factory", [microsoft_word_updater, wps_writer_updater], ids=["word", "wps"]
)
def test_update_fields_fills_toc_and_reports_pages(tmp_path: Path, factory) -> None:
    updater = _updater_or_skip(factory)
    source = build_multipage_document(tmp_path / "input.docx")
    before = document_text(source)
    assert before.count(HEADINGS[0]) == 1, "更新前目录应当还是空的"

    outcome = updater.update_fields(source, tmp_path / "output.docx")

    assert outcome.output_path.is_file()
    assert outcome.page_count is not None and outcome.page_count >= 3
    after = document_text(outcome.output_path)
    for title in HEADINGS:
        assert after.count(title) >= 2, f"目录里没有出现 {title}"


@pytest.mark.parametrize(
    "factory", [microsoft_word_updater, wps_writer_updater], ids=["word", "wps"]
)
def test_many_sections_with_headers_finish_and_get_real_page_numbers(
    tmp_path: Path, factory
) -> None:
    """多节带页眉页脚的文档必须刷得完，而且每节页脚拿到的是真页码。

    这条守的是 §19 的一个真实故障：早先按 NextStoryRange 串故事区，本机的 WPS 在
    页眉页脚那条链上不收敛，同一个页眉无限重复，21MB 的报告 900 秒超时都刷不完，
    改成按节取页眉页脚后 8 秒完成。单节文档没有那条链，测不出来。
    """
    updater = _updater_or_skip(factory)
    updater.timeout_seconds = 120
    source = build_sectioned_document(tmp_path / "input.docx")

    outcome = updater.update_fields(source, tmp_path / "output.docx")

    assert outcome.page_count is not None and outcome.page_count >= 7
    # 每节各占一页，页脚里的 PAGE 应当逐节递增；刷不到页脚时它们会全是初始的 1。
    numbers = footer_text(outcome.output_path)
    assert len(set(numbers)) == len(numbers), f"各节页脚一模一样，页码没算：{numbers}"


@pytest.mark.parametrize(
    "factory", [microsoft_word_updater, wps_writer_updater], ids=["word", "wps"]
)
def test_input_document_is_never_modified(tmp_path: Path, factory) -> None:
    """Builder 的产物必须原封不动——失败时要能用同一份输入重试。"""
    updater = _updater_or_skip(factory)
    source = build_multipage_document(tmp_path / "input.docx")
    original = source.read_bytes()

    updater.update_fields(source, tmp_path / "output.docx")

    assert source.read_bytes() == original
    assert not list(tmp_path.glob("~$*")), "输入目录不应残留 Word owner file"


def test_missing_input_fails_with_stage(tmp_path: Path) -> None:
    updater = microsoft_word_updater(timeout_seconds=60)

    with pytest.raises(FieldUpdateError) as excinfo:
        updater.update_fields(tmp_path / "nope.docx", tmp_path / "out.docx")

    assert excinfo.value.code == "REPORT_FIELD_UPDATE_FAILED"
    assert excinfo.value.stage == "open"


def test_concurrent_updates_never_run_in_word_at_the_same_time(tmp_path: Path) -> None:
    """设计 §25.3：两个并发请求必须串行进 Word，且等待状态可见。

    这条同时守着超时清理的正确性——`_run_isolated` 靠"动手前后的进程差集"认领自己
    起的 Word，时间窗一旦重叠就会杀错人。队列被拿掉时这个用例会开始间歇失败。
    """
    updater = _updater_or_skip(microsoft_word_updater)
    source = build_multipage_document(tmp_path / "input.docx")
    spans: dict[str, tuple[float, float]] = {}
    waits: dict[str, int] = {}
    failures: list[BaseException] = []

    def run(name: str) -> None:
        try:
            started = time.monotonic()
            outcome = updater.update_fields(
                source,
                tmp_path / f"out-{name}.docx",
                on_wait=lambda ahead, n=name: waits.__setitem__(n, ahead),
            )
            # 两端都由 outcome 推出来，避免把"函数返回"当成"Word 结束"——锁在函数
            # 内部就释放了，用返回时刻会把后一个的开始算成早于前一个的结束。
            word_start = started + outcome.queued_seconds
            spans[name] = (word_start, word_start + outcome.elapsed_seconds)
        except BaseException as exc:  # noqa: BLE001 - 线程内异常要带回主线程断言
            failures.append(exc)

    threads = [threading.Thread(target=run, args=(name,)) for name in ("a", "b")]
    for thread in threads:
        thread.start()
    for thread in threads:
        thread.join(600)

    assert not failures, failures
    assert set(spans) == {"a", "b"}

    (a_in, a_out), (b_in, b_out) = spans["a"], spans["b"]
    assert a_out <= b_in or b_out <= a_in, f"两次更新在 Word 里重叠了：{spans}"
    assert waits, "两个请求都没有排队，队列没有生效"


def test_timeout_kills_only_our_instance_and_removes_output(tmp_path: Path) -> None:
    """超时清理不能误伤用户已经开着的 Word/WPS，且不留下半成品。"""
    updater = _updater_or_skip(microsoft_word_updater)
    updater.timeout_seconds = 0.001
    source = build_multipage_document(tmp_path / "input.docx")
    output = tmp_path / "output.docx"

    with pytest.raises(FieldUpdateError) as excinfo:
        updater.update_fields(source, output)

    assert excinfo.value.stage == "timeout"
    assert not output.exists()
