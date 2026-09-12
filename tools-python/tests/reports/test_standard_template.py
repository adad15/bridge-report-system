"""校验随系统发布的首个标准模板。

测的是签入仓库的产物本身（templates/report/periodic_inspection_v1/），不是生成它的
脚本——模板也可能被人直接用 Word 改，那时这些用例同样要拦住不合契约的改动。
生成脚本在 scripts/report-template/，不属于运行时代码。
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn

from bridge_report_tools.reports.contract import (
    FIELDS_EXPLICITLY_REJECTED,
    PERIODIC_INSPECTION_V1,
    personnel_placeholder,
)
from bridge_report_tools.reports.docx_builder import build_report
from bridge_report_tools.reports.docx_scan import scan_template
from bridge_report_tools.reports.errors import ReportBuildError
from bridge_report_tools.reports.styles import (
    REQUIRED_BUILDER_STYLES,
    STYLE_CARD_BAND,
    STYLE_CARD_CELL,
    STYLE_CARD_HEADER,
    STYLE_PHOTO,
)
from bridge_report_tools.reports.template_validator import (
    TemplateConfig,
    validate_template,
)

from tests.reports.context_fixtures import SCALARS, defect_row, part, photo, png
from tests.reports.context_fixtures import context as report_context


REPO_ROOT = Path(__file__).resolve().parents[3]
TEMPLATE_DIR = REPO_ROOT / "templates" / "report" / "periodic_inspection_v1"
TEMPLATE_DOCX = TEMPLATE_DIR / "periodic-inspection-v1.docx"
TEMPLATE_CONFIG = TEMPLATE_DIR / "template.json"

EXPECTED_ANCHORS = [
    "PERSONNEL_TABLE",
    "BRIDGE_PROFILE",
    "EQUIPMENT_LIST",
    "DEFECT_TABLES:SUPERSTRUCTURE",
    "DEFECT_PHOTOS:SUPERSTRUCTURE",
    "PREVIOUS_COMPARISON:SUPERSTRUCTURE",
    "DEFECT_TABLES:SUBSTRUCTURE",
    "DEFECT_PHOTOS:SUBSTRUCTURE",
    "PREVIOUS_COMPARISON:SUBSTRUCTURE",
    "DEFECT_TABLES:DECK",
    "DEFECT_PHOTOS:DECK",
    "PREVIOUS_COMPARISON:DECK",
    "COMPONENT_WEIGHTS",
    "ASSESSMENT_RESULT",
    "CONTROL_INDICATOR",
    "OVERALL_ASSESSMENT",
    "CONCLUSION",
    "ASSESSMENT_APPENDIX",
    "BRIDGE_CARD",
]

pytestmark = pytest.mark.skipif(
    not TEMPLATE_DOCX.is_file(),
    reason=f"未找到已发布的模板 {TEMPLATE_DOCX}；用 scripts/report-template 重新生成",
)


@pytest.fixture(scope="module")
def config() -> TemplateConfig:
    payload = json.loads(TEMPLATE_CONFIG.read_text(encoding="utf-8"))
    return TemplateConfig(
        table_number_formats=payload["table_number_formats"],
        required_personnel_roles=payload["required_personnel_roles"],
    )


def test_shipped_template_passes_its_own_contract(config: TemplateConfig) -> None:
    result = validate_template(TEMPLATE_DOCX, config)

    assert result.is_valid, result.codes()


def test_shipped_config_declares_the_expected_contract_type() -> None:
    payload = json.loads(TEMPLATE_CONFIG.read_text(encoding="utf-8"))

    assert payload["contract_type"] == PERIODIC_INSPECTION_V1.contract_type


def test_anchor_order_follows_the_official_report_structure(config: TemplateConfig) -> None:
    """第 2 章按上部、下部、桥面系各起一节，与两份正式报告一致。"""
    result = validate_template(TEMPLATE_DOCX, config)

    assert result.anchors_in_document_order == EXPECTED_ANCHORS


def test_template_uses_only_toc_and_page_fields() -> None:
    """两份原件里有 913 个 STYLEREF、458 个 SEQ、455 个 REF；新模板一个都不许有。"""
    keywords = {hit.keyword for hit in scan_template(TEMPLATE_DOCX).fields}

    assert keywords <= {"TOC", "PAGE", "NUMPAGES"}
    assert not keywords & FIELDS_EXPLICITLY_REJECTED


def test_every_placeholder_is_known_to_the_contract() -> None:
    scan = scan_template(TEMPLATE_DOCX)

    assert {hit.name for hit in scan.placeholders} <= PERIODIC_INSPECTION_V1.all_placeholders
    assert all(hit.is_well_formed for hit in scan.placeholders)
    assert scan.unbalanced_braces == []


def test_config_covers_every_numbered_anchor() -> None:
    """不给配置时每个需要编号的内容块都应被点名，说明 template.json 没漏项。"""
    payload = json.loads(TEMPLATE_CONFIG.read_text(encoding="utf-8"))
    result = validate_template(TEMPLATE_DOCX, TemplateConfig())

    assert result.codes().count("template_number_format_missing") == len(
        payload["table_number_formats"]
    )


def test_body_and_appendix_sections_have_running_header_and_footer() -> None:
    sections = Document(str(TEMPLATE_DOCX)).sections

    # 封面、声明页、签字页、目录、正文、附录1、附录2。
    assert len(sections) == 7
    # 前四节不带页眉页脚。
    for section in sections[:4]:
        assert all(not p.text.strip() for p in section.header.paragraphs)
    for section in sections[4:]:
        assert "定期检测报告" in section.header.paragraphs[0].text


def test_only_the_last_appendix_section_is_landscape() -> None:
    """附录1 是竖版的评定卡片，只有附录2（基本状况卡片）走横版。"""
    sections = Document(str(TEMPLATE_DOCX)).sections

    for section in sections[:-1]:
        assert section.orientation == WD_ORIENT.PORTRAIT
        assert section.page_width < section.page_height
    landscape = sections[-1]
    assert landscape.orientation == WD_ORIENT.LANDSCAPE
    assert landscape.page_width > landscape.page_height


def test_appendix_one_shares_the_body_page_setup() -> None:
    """附录1 的页面设置照抄正文节——同样的纸张和页边距，卡片才排得下。"""
    sections = Document(str(TEMPLATE_DOCX)).sections
    body, appendix_one = sections[4], sections[5]

    assert appendix_one.page_width == body.page_width
    assert appendix_one.left_margin == body.left_margin
    assert appendix_one.right_margin == body.right_margin


def test_builder_styles_exist() -> None:
    """Docx Builder 只许用命名样式，不在代码里散落字号边距（设计 §18）。"""
    names = {style.name for style in Document(str(TEMPLATE_DOCX)).styles}

    assert REQUIRED_BUILDER_STYLES <= names


def test_card_styles_are_small_song_with_roman_digits() -> None:
    """两张附录卡片用宋体小五，数字 Times New Roman。

    正文表格的五号在卡片里排不下：附录1 一行三对「标签/取值」，附录2 一行三组
    「编号/名称/取值」还带联系电话，字大一号就到处折行。
    """
    styles = Document(str(TEMPLATE_DOCX)).styles

    for name in (STYLE_CARD_HEADER, STYLE_CARD_CELL, STYLE_CARD_BAND):
        style = styles[name]
        fonts = style.element.rPr.find(qn("w:rFonts"))
        assert fonts.get(qn("w:eastAsia")) == "宋体", name
        assert fonts.get(qn("w:ascii")) == "Times New Roman", name
        assert style.font.size.pt == 9, name


def test_photo_style_keeps_image_with_caption() -> None:
    """图片与图题不可跨页拆分（设计 §11.5）。"""
    style = Document(str(TEMPLATE_DOCX)).styles[STYLE_PHOTO]

    assert style.paragraph_format.keep_with_next is True


def test_every_placeholder_has_a_value_source() -> None:
    """模板用到的每个占位符，ReportContext 都得给得出取值。

    模板校验只保证"名字认识"；名字认识但上下文不提供，生成出来就是个空。
    """
    scan = scan_template(TEMPLATE_DOCX)
    declared_roles = json.loads(TEMPLATE_CONFIG.read_text(encoding="utf-8"))[
        "required_personnel_roles"
    ]
    available = set(SCALARS) | {personnel_placeholder(role) for role in declared_roles}

    assert {hit.name for hit in scan.placeholders} <= available


def test_shipped_template_builds_a_complete_report(tmp_path: Path) -> None:
    """发布模板 + 一份完整上下文 = 一份没有残留、没有缺块的报告。

    这是端到端的那道底线：模板里的锚点全部装配完，正文里既不剩 [[REPORT:...]]，
    也不剩 {{占位符}}——最终校验（设计 §20.2、§20.3）查的就是这两条。
    """
    archive = tmp_path / "archive"
    (archive / "photos").mkdir(parents=True)
    png(archive / "photos/a.png", 800, 600)
    photos = [photo("照片2.1-1", "photos/a.png"), photo("照片2.1-2", "photos/a.png")]
    ctx = report_context(
        parts=[
            part(
                "SUPERSTRUCTURE",
                "上部结构",
                rows=[defect_row(1, photo_numbers=["照片2.1-1", "照片2.1-2"])],
                photos=photos,
            ),
            part("SUBSTRUCTURE", "下部结构", rows=[defect_row(1)], photos=[]),
        ],
        personnel=[
            {"full_name": "张三", "role_code": "approver", "organization": "某某院",
             "professional_title": "教高", "qualification_certificate_no": "JC-001"},
        ],
        equipment=[
            {"equipment_name": "裂缝观测仪", "model_spec": "ZBL-F130",
             "asset_number": "SB-01", "measurement_range": "0-6mm", "accuracy": "0.01mm",
             "calibration_certificate_no": "JD-1",
             "calibration_valid_until": "2027-01-31", "purpose": "裂缝宽度"},
        ],
    )

    result = build_report(TEMPLATE_DOCX, ctx, tmp_path / "report.docx", archive)

    assert result.blocks_rendered == EXPECTED_ANCHORS
    assert result.missing_placeholders == []
    document = Document(str(result.output_path))
    text = "\n".join(
        [paragraph.text for paragraph in document.paragraphs]
        + [
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        ]
    )
    assert "[[REPORT:" not in text
    assert "{{" not in text
