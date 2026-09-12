"""表头样表（设计 §7.7）。

模板作者在 Word 里画一张只有表头行的表，拖出想要的列宽、写好表头文字；生成器照它
建表，行数和合并仍由数据决定。

这里守的核心是一条：**样表的列数必须与契约一致**。模板作者删掉一列或调换两列，
生成器照填不误——「病害位置」的内容会印到「病害类型」下面，看着完全正常但全错。
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.shared import Cm

from bridge_report_tools.reports.contract import DEFECT_TABLE_COLUMNS, table_marker
from bridge_report_tools.reports.docx_builder import build_report
from bridge_report_tools.reports.docx_scan import scan_table_samples
from bridge_report_tools.reports.template_validator import validate_template

from tests.reports.context_fixtures import context, defect_row, part, png
from tests.reports.template_fixtures import (
    CORE_ANCHORS,
    build_builder_template,
    build_template,
    valid_config,
)


#: 病害表的十列宽度，合计 16cm（正文版心）。
DEFECT_WIDTHS = [1.0, 2.2, 1.8, 1.8, 1.8, 3.4, 0.8, 1.2, 1.2, 0.8]


def add_sample_table(document, headers, style=None):
    table = document.add_table(1, len(headers), style=style)
    for index, header in enumerate(headers):
        table.cell(0, index).text = header
    return table


def add_sample(document, block: str, headers, widths_cm=None, style=None) -> None:
    """在文档末尾加一个"标记段 + 表头样表"。

    样表用什么表格样式不影响扫描（只读表头文字和列宽），所以默认不指定——
    最小合规模板里没有「报告表格」这个样式。
    """
    document.add_paragraph(table_marker(block))
    table = add_sample_table(document, headers, style)
    if widths_cm:
        table.autofit = False
        for index, width in enumerate(widths_cm):
            table.columns[index].width = Cm(width)
            for row in table.rows:
                row.cells[index].width = Cm(width)


def superstructure():
    return context(parts=[part("SUPERSTRUCTURE", "上部结构",
                               rows=[defect_row(1)], photos=[])])


@pytest.fixture
def archive(tmp_path: Path) -> Path:
    root = tmp_path / "archive"
    (root / "photos").mkdir(parents=True)
    png(root / "photos/a.png", 800, 600)
    return root


# --------------------------------------------------------------------------
# 扫描
# --------------------------------------------------------------------------


def test_scan_reads_headers_and_widths(tmp_path: Path) -> None:
    document = Document()
    add_sample(document, "DEFECT_TABLES", DEFECT_TABLE_COLUMNS, DEFECT_WIDTHS)
    path = tmp_path / "t.docx"
    document.save(str(path))

    samples = scan_table_samples(Document(str(path)))

    assert len(samples) == 1
    sample = samples[0]
    assert sample.block == "DEFECT_TABLES"
    assert sample.has_table
    assert sample.headers == list(DEFECT_TABLE_COLUMNS)
    assert [round(width / 360000, 1) for width in sample.widths] == DEFECT_WIDTHS


def test_scan_notices_a_marker_without_a_table(tmp_path: Path) -> None:
    document = Document()
    document.add_paragraph(table_marker("DEFECT_TABLES"))
    document.add_paragraph("后面不是表格")
    path = tmp_path / "t.docx"
    document.save(str(path))

    assert scan_table_samples(Document(str(path)))[0].has_table is False


# --------------------------------------------------------------------------
# 校验（设计 §7.7）
# --------------------------------------------------------------------------


def test_template_without_samples_is_still_valid(tmp_path: Path) -> None:
    """样表是可选的：没画就退回生成器的默认列宽，老模板照常能用。"""
    result = validate_template(build_template(tmp_path / "t.docx"), valid_config())

    assert result.is_valid, result.codes()


def test_sample_with_the_right_columns_is_valid(tmp_path: Path) -> None:
    def customize(document):
        add_sample(document, "DEFECT_TABLES", DEFECT_TABLE_COLUMNS, DEFECT_WIDTHS)

    result = validate_template(
        build_template(tmp_path / "t.docx", customize=customize), valid_config()
    )

    assert result.is_valid, result.codes()


def test_sample_with_a_missing_column_is_rejected(tmp_path: Path) -> None:
    """少一列就会把内容印到错误的列下面——看着正常但全错，必须拦住。"""

    def customize(document):
        add_sample(document, "DEFECT_TABLES", DEFECT_TABLE_COLUMNS[:-1])

    result = validate_template(
        build_template(tmp_path / "t.docx", customize=customize), valid_config()
    )

    assert not result.is_valid
    assert "template_table_sample_column_mismatch" in result.codes()


def test_sample_marker_without_a_table_is_rejected(tmp_path: Path) -> None:
    def customize(document):
        document.add_paragraph(table_marker("DEFECT_TABLES"))

    result = validate_template(
        build_template(tmp_path / "t.docx", customize=customize), valid_config()
    )

    assert not result.is_valid
    assert "template_table_sample_missing_table" in result.codes()


def test_sample_for_an_unknown_block_is_rejected(tmp_path: Path) -> None:
    def customize(document):
        add_sample(document, "CONCLUSION", ("一", "二"))

    result = validate_template(
        build_template(tmp_path / "t.docx", customize=customize), valid_config()
    )

    assert not result.is_valid
    assert "template_table_sample_unknown_block" in result.codes()


def test_duplicate_sample_is_rejected(tmp_path: Path) -> None:
    def customize(document):
        add_sample(document, "DEFECT_TABLES", DEFECT_TABLE_COLUMNS)
        add_sample(document, "DEFECT_TABLES", DEFECT_TABLE_COLUMNS)

    result = validate_template(
        build_template(tmp_path / "t.docx", customize=customize), valid_config()
    )

    assert not result.is_valid
    assert "template_table_sample_duplicated" in result.codes()


def test_sample_marker_may_carry_a_hint_after_it(tmp_path: Path) -> None:
    """样表整段删除，所以标记后面可以跟一句给模板作者看的说明。"""

    def customize(document):
        document.add_paragraph(f"{table_marker('DEFECT_TABLES')}  病害检查表：拖列宽")
        add_sample_table(document, DEFECT_TABLE_COLUMNS)

    result = validate_template(
        build_template(tmp_path / "t.docx", customize=customize), valid_config()
    )

    assert result.is_valid, result.codes()


def test_sample_marker_mid_sentence_is_rejected(tmp_path: Path) -> None:
    """写在句子中间，删段会把前面的话一起删了。"""

    def customize(document):
        document.add_paragraph(f"见 {table_marker('DEFECT_TABLES')} 样表")
        add_sample_table(document, DEFECT_TABLE_COLUMNS)

    result = validate_template(
        build_template(tmp_path / "t.docx", customize=customize), valid_config()
    )

    assert not result.is_valid
    assert "template_table_sample_not_leading" in result.codes()


# --------------------------------------------------------------------------
# 装配
# --------------------------------------------------------------------------


def test_builder_uses_the_sample_widths_and_headers(
    tmp_path: Path, archive: Path
) -> None:
    def customize(document):
        add_sample(
            document,
            "DEFECT_TABLES",
            ("序号", "部件", "编号", "位置", "类型", "特征", "标度", "扣分", "评分", "照片"),
            DEFECT_WIDTHS,
        )

    template = build_builder_template(
        tmp_path / "template.docx",
        anchors=("DEFECT_TABLES:SUPERSTRUCTURE",),
        customize=customize,
    )
    result = build_report(template, superstructure(), tmp_path / "out.docx", archive)

    table = Document(str(result.output_path)).tables[0]
    # 表头文字来自样表，不是代码里的默认值。
    assert [cell.text for cell in table.rows[0].cells][:3] == ["序号", "部件", "编号"]
    # 列宽也来自样表。
    assert [round(column.width / 360000, 1) for column in table.columns] == DEFECT_WIDTHS


def test_sample_area_never_reaches_the_report(tmp_path: Path, archive: Path) -> None:
    """样表是调版式用的，不是报告内容；留在交付文件里就是一张空表加一行标记。"""

    def customize(document):
        add_sample(document, "DEFECT_TABLES", DEFECT_TABLE_COLUMNS, DEFECT_WIDTHS)

    template = build_builder_template(
        tmp_path / "template.docx",
        anchors=("DEFECT_TABLES:SUPERSTRUCTURE",),
        customize=customize,
    )
    result = build_report(template, superstructure(), tmp_path / "out.docx", archive)

    document = Document(str(result.output_path))
    assert all("[[TABLE:" not in paragraph.text for paragraph in document.paragraphs)
    # 只剩病害表本身，样表那张空表没了。
    assert len(document.tables) == 1
    assert len(document.tables[0].rows) == 2


def test_removing_a_sample_never_eats_a_section_break(
    tmp_path: Path, archive: Path
) -> None:
    """带分节符的段落除了 sectPr 什么都没有，看着也是"空段"。

    删样表时连它一起收走，正文就失去自己的页面设置、继承最后那一节——实测报告从
    104 页涨到 166 页，因为正文继承了横向节，页高从 29.7cm 变成 21cm。
    """

    def customize(document):
        # 造第二节，然后把样表塞进第一节的末尾（分节符段落之前）。
        document.add_section(WD_SECTION.NEW_PAGE)
        document.add_paragraph("第二节的内容")
        add_sample(document, "DEFECT_TABLES", DEFECT_TABLE_COLUMNS, DEFECT_WIDTHS)
        body = document.element.body
        break_paragraph = next(
            child
            for child in body.iterchildren()
            if child.tag == qn("w:p") and child.xpath("./w:pPr/w:sectPr")
        )
        # 样表标记、样表、以及后面那个空段，全搬到分节符之前。
        for element in list(body.iterchildren())[-3:]:
            break_paragraph.addprevious(element)

    template = build_builder_template(
        tmp_path / "template.docx",
        anchors=("DEFECT_TABLES:SUPERSTRUCTURE",),
        customize=customize,
    )
    before = len(Document(str(template)).sections)
    result = build_report(template, superstructure(), tmp_path / "out.docx", archive)

    assert len(Document(str(result.output_path)).sections) == before


def test_builder_falls_back_without_a_sample(tmp_path: Path, archive: Path) -> None:
    template = build_builder_template(
        tmp_path / "template.docx", anchors=("DEFECT_TABLES:SUPERSTRUCTURE",)
    )
    result = build_report(template, superstructure(), tmp_path / "out.docx", archive)

    table = Document(str(result.output_path)).tables[0]
    assert [cell.text for cell in table.rows[0].cells] == list(DEFECT_TABLE_COLUMNS)


def test_shipped_template_anchor_order_is_unaffected_by_samples(tmp_path: Path) -> None:
    """样表标记不是内容锚点，不该混进"锚点及文档顺序"那份清单里。"""

    def customize(document):
        add_sample(document, "DEFECT_TABLES", DEFECT_TABLE_COLUMNS)

    result = validate_template(
        build_template(tmp_path / "t.docx", customize=customize), valid_config()
    )

    assert result.anchors_in_document_order == list(CORE_ANCHORS)
