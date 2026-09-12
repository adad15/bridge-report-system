"""Docx Builder 的装配用例（设计 §10、§11、§18）。

守两类事：

* **不出错误内容**：病害表的照片编号列与图题必须完全一致（§11.8 的致命失败模式）；
  没病害不出空表，没照片不出占位图；装不出来的块报错而不是留个空段落。
* **不破坏模板**：模板文件本身不被改；样式、分节和页眉页脚原样留着。
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from bridge_report_tools.reports.conclusion import (
    COMPARISON_CAVEAT,
    MAINTENANCE_ADVICE,
    NO_COMPARISON_TEXT,
    NO_CONTROL_INDICATOR_TEXT,
)
from bridge_report_tools.reports.contract import PERIODIC_INSPECTION_V1
from bridge_report_tools.reports.docx_builder import (
    DEFECT_TABLE_HEADERS,
    NO_DEFECT_TEXT,
    PENDING_BLOCKS,
    RENDERERS,
    RESULT_TABLE_HEADERS,
    WEIGHT_TABLE_HEADERS,
    build_report,
)
from bridge_report_tools.reports.errors import ReportBuildError
from bridge_report_tools.reports.report_context import ReportContext
from bridge_report_tools.reports.styles import (
    STYLE_CARD_BAND,
    STYLE_CARD_CELL,
    STYLE_CARD_HEADER,
    STYLE_PHOTO_CAPTION,
)

from tests.reports.context_fixtures import (
    NO_ASSESSMENT,
    assessment,
    comparison,
    context,
    defect_row,
    part,
    photo,
    png,
)
from tests.reports.template_fixtures import build_builder_template


ARCHIVE = "photos/a.png"


@pytest.fixture
def archive(tmp_path: Path) -> Path:
    root = tmp_path / "archive"
    (root / "photos").mkdir(parents=True)
    png(root / ARCHIVE, 800, 600)
    return root


def build(tmp_path: Path, archive_root: Path, ctx: ReportContext, **kwargs):
    template = build_builder_template(tmp_path / "template.docx", **kwargs)
    return (
        build_report(template, ctx, tmp_path / "out.docx", archive_root),
        template,
    )


def superstructure(rows: list[dict], photos: list[dict]) -> ReportContext:
    return context(parts=[part("SUPERSTRUCTURE", "上部结构", rows=rows, photos=photos)])


def body_text(path: Path) -> list[str]:
    document = Document(str(path))
    return [paragraph.text for paragraph in document.paragraphs]


def tables(path: Path):
    return Document(str(path)).tables


# --------------------------------------------------------------------------
# 标量占位符
# --------------------------------------------------------------------------


def test_scalars_are_replaced(tmp_path: Path, archive: Path) -> None:
    result, _ = build(tmp_path, archive, superstructure([], []))

    assert "Q2026-1" in body_text(result.output_path)
    assert "百股大桥" in body_text(result.output_path)
    assert result.missing_placeholders == []


def test_placeholder_split_across_runs_is_replaced(tmp_path: Path, archive: Path) -> None:
    """Word 会把 {{report_no}} 拆进多个 run，拆了照样要替换掉。"""

    def customize(document) -> None:
        paragraph = document.add_paragraph()
        for chunk in ("{{repo", "rt_", "no}}"):
            paragraph.add_run(chunk)

    result, _ = build(
        tmp_path, archive, superstructure([], []), anchors=(), customize=customize
    )

    assert "Q2026-1" in body_text(result.output_path)
    assert "{{" not in "".join(body_text(result.output_path))


def test_run_styles_survive_replacement(tmp_path: Path, archive: Path) -> None:
    """段落里只有占位符那一段被换掉，前后 run 的样式不能被抹平。"""

    def customize(document) -> None:
        paragraph = document.add_paragraph()
        paragraph.add_run("编号：").bold = True
        paragraph.add_run("{{report_no}}")
        paragraph.add_run("（终稿）").italic = True

    result, _ = build(
        tmp_path, archive, superstructure([], []), anchors=(), customize=customize
    )

    paragraph = [p for p in Document(str(result.output_path)).paragraphs if "编号：" in p.text][0]
    assert paragraph.text == "编号：Q2026-1（终稿）"
    assert paragraph.runs[0].bold is True
    assert paragraph.runs[-1].italic is True


def test_placeholders_in_header_and_footer_are_replaced(
    tmp_path: Path, archive: Path
) -> None:
    """动态页眉是被禁的（§7.6），但页眉里放标量占位符是允许的写法。"""

    def customize(document) -> None:
        section = document.sections[0]
        section.header.is_linked_to_previous = False
        section.header.paragraphs[0].text = "{{bridge_name}}定期检测报告"
        section.footer.is_linked_to_previous = False
        section.footer.paragraphs[0].text = "{{report_no}}"

    result, _ = build(
        tmp_path, archive, superstructure([], []), anchors=(), customize=customize
    )

    section = Document(str(result.output_path)).sections[0]
    assert section.header.paragraphs[0].text == "百股大桥定期检测报告"
    assert section.footer.paragraphs[0].text == "Q2026-1"


def test_personnel_scalar_groups_names_by_role(tmp_path: Path, archive: Path) -> None:
    ctx = context(
        personnel=[
            {"full_name": "张三", "role_code": "approver", "organization": "某某院",
             "professional_title": None, "qualification_certificate_no": None},
            {"full_name": "李四", "role_code": "approver", "organization": None,
             "professional_title": None, "qualification_certificate_no": None},
            {"full_name": "王五", "role_code": "reviewer", "organization": None,
             "professional_title": None, "qualification_certificate_no": None},
        ]
    )
    result, _ = build(
        tmp_path,
        archive,
        ctx,
        anchors=(),
        placeholders=("personnel.approver.names", "personnel.reviewer.names"),
    )

    assert "张三、李四" in body_text(result.output_path)
    assert "王五" in body_text(result.output_path)


# --------------------------------------------------------------------------
# 病害表（设计 §10）
# --------------------------------------------------------------------------


def test_defect_table_has_the_official_columns(tmp_path: Path, archive: Path) -> None:
    ctx = superstructure([defect_row(1)], [])
    result, _ = build(tmp_path, archive, ctx)

    table = tables(result.output_path)[0]
    assert [cell.text for cell in table.rows[0].cells] == list(DEFECT_TABLE_HEADERS)
    assert table.rows[1].cells[2].text == "1-1#板"


def test_defect_table_rows_do_not_chain_to_each_other(
    tmp_path: Path, archive: Path
) -> None:
    """表格紧跟前面的文字，不整体另起一页。

    单元格段落若都设了"与下段同页"，一行接一行连成一串，整张表就变成不可分割的块，
    Word 只好把它整体推到下一页——实测 50 行的下部结构病害表因此与前面的概要文字
    断开。行内不断页应当用 cantSplit 表达，而不是段落的 keepNext。
    """
    ctx = superstructure([defect_row(n) for n in range(1, 6)], [])
    result, _ = build(tmp_path, archive, ctx)

    table = tables(result.output_path)[0]
    keep_next = [
        paragraph.paragraph_format.keep_with_next
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
    ]
    assert not any(keep_next)
    # 行内不跨页断开仍然要有，并且 cantSplit 必须排在 tblHeader 之前。
    header_properties = [child.tag for child in table.rows[0]._tr.find(qn("w:trPr"))]
    assert header_properties == [qn("w:cantSplit"), qn("w:tblHeader")]
    for row in table.rows[1:]:
        assert row._tr.find(qn("w:trPr")).find(qn("w:cantSplit")) is not None


def test_defect_table_caption_uses_the_template_number_format(
    tmp_path: Path, archive: Path
) -> None:
    ctx = superstructure([defect_row(1)], [])
    result, _ = build(tmp_path, archive, ctx)

    assert any(text.startswith("表2.1-1") for text in body_text(result.output_path))


def test_part_without_defects_gets_a_sentence_not_an_empty_table(
    tmp_path: Path, archive: Path
) -> None:
    """不生成空表格行，也不声称构件完好或不存在（设计 §10.3）。"""
    result, _ = build(tmp_path, archive, superstructure([], []))

    assert NO_DEFECT_TEXT in body_text(result.output_path)
    assert tables(result.output_path) == []


def test_defect_without_photos_gets_no_fake_placeholder(
    tmp_path: Path, archive: Path
) -> None:
    ctx = superstructure([defect_row(1)], [])
    result, _ = build(tmp_path, archive, ctx)

    table = tables(result.output_path)[0]
    assert table.rows[1].cells[-1].text == "—"
    assert _picture_count(result.output_path) == 0


# --------------------------------------------------------------------------
# 照片（设计 §11）
# --------------------------------------------------------------------------


def test_photo_numbers_in_table_and_captions_are_identical(
    tmp_path: Path, archive: Path
) -> None:
    """本方案唯一的致命失败模式（设计 §11.8）：两处号码对不上就按号找不到图。"""
    photos = [photo("照片2.1-1", ARCHIVE), photo("照片2.1-2", ARCHIVE)]
    ctx = superstructure(
        [defect_row(1, photo_numbers=["照片2.1-1"]), defect_row(2, photo_numbers=["照片2.1-2"])],
        photos,
    )
    result, _ = build(tmp_path, archive, ctx)

    document = Document(str(result.output_path))
    from_table = [row.cells[-1].text for row in document.tables[0].rows[1:]]
    from_captions = [
        paragraph.text.split("  ")[0]
        for table in document.tables[1:]
        for row in table.rows
        for cell in row.cells
        for paragraph in cell.paragraphs
        if paragraph.style.name == STYLE_PHOTO_CAPTION and paragraph.text
    ]
    assert from_table == ["照片2.1-1", "照片2.1-2"]
    assert from_captions == from_table


def test_photos_are_laid_out_two_per_row(tmp_path: Path, archive: Path) -> None:
    photos = [photo(f"照片2.1-{n}", ARCHIVE) for n in range(1, 4)]
    ctx = superstructure(
        [defect_row(1, photo_numbers=[p["report_number"] for p in photos])], photos
    )
    result, _ = build(tmp_path, archive, ctx)

    layout = tables(result.output_path)[1]
    # 整块照片一张表，每两行一组（图片行 + 图题行）：3 张 -> 2 组 -> 4 行。
    assert (len(layout.rows), len(layout.columns)) == (4, 2)


def test_photo_block_is_one_table_so_word_cannot_merge_it(
    tmp_path: Path, archive: Path
) -> None:
    """紧挨着的两张表会被 Word 合并——实测 6 张读回来只剩 5 张。

    照片块做成一张表就没有内部相邻；块与块之间由 BlockInserter 垫空段落隔开。
    """
    photos = [photo(f"照片2.1-{n}", ARCHIVE) for n in range(1, 5)]
    ctx = superstructure(
        [defect_row(1, photo_numbers=[p["report_number"] for p in photos])], photos
    )
    result, _ = build(tmp_path, archive, ctx)

    document = Document(str(result.output_path))
    body = list(document.element.body.iterchildren())
    adjacent = [
        index
        for index in range(len(body) - 1)
        if body[index].tag == qn("w:tbl") and body[index + 1].tag == qn("w:tbl")
    ]
    assert adjacent == []


def test_odd_last_photo_stays_in_the_left_column(tmp_path: Path, archive: Path) -> None:
    """最后一行只剩一张时右栏留空，不把它拉宽（设计 §11.6）。"""
    photos = [photo(f"照片2.1-{n}", ARCHIVE) for n in range(1, 4)]
    ctx = superstructure(
        [defect_row(1, photo_numbers=[p["report_number"] for p in photos])], photos
    )
    result, _ = build(tmp_path, archive, ctx)

    layout = tables(result.output_path)[1]
    assert layout.cell(3, 0).text == "照片2.1-3  横向裂缝"
    assert layout.cell(3, 1).text == ""
    assert layout.columns[0].width == layout.columns[1].width


def test_photo_is_scaled_into_the_frame_without_stretching(
    tmp_path: Path, archive: Path
) -> None:
    png(archive / "photos/wide.png", 1600, 400)
    photos = [photo("照片2.1-1", "photos/wide.png")]
    ctx = superstructure([defect_row(1, photo_numbers=["照片2.1-1"])], photos)
    result, _ = build(tmp_path, archive, ctx)

    extent = _first_picture_extent(result.output_path)
    assert extent is not None
    width, height = extent
    assert abs(width / height - 4.0) < 0.01  # 1600x400 的比例原样保留


def test_caption_without_title_is_just_the_number(tmp_path: Path, archive: Path) -> None:
    photos = [photo("照片2.1-1", ARCHIVE, title=None)]
    ctx = superstructure([defect_row(1, photo_numbers=["照片2.1-1"])], photos)
    result, _ = build(tmp_path, archive, ctx)

    assert tables(result.output_path)[1].cell(1, 0).text == "照片2.1-1"


def test_missing_photo_file_aborts_the_build(tmp_path: Path, archive: Path) -> None:
    """归档文件读不到就中止，绝不出一份少图的报告（设计 §11.10）。"""
    photos = [photo("照片2.1-1", "photos/gone.png")]
    ctx = superstructure([defect_row(1, photo_numbers=["照片2.1-1"])], photos)

    with pytest.raises(ReportBuildError) as raised:
        build(tmp_path, archive, ctx)

    assert raised.value.code == "report_photo_file_missing"


def test_unreadable_photo_file_aborts_the_build(tmp_path: Path, archive: Path) -> None:
    (archive / "photos/broken.png").write_bytes(b"not an image")
    photos = [photo("照片2.1-1", "photos/broken.png")]
    ctx = superstructure([defect_row(1, photo_numbers=["照片2.1-1"])], photos)

    with pytest.raises(ReportBuildError) as raised:
        build(tmp_path, archive, ctx)

    assert raised.value.code == "report_photo_unreadable"


# --------------------------------------------------------------------------
# 人员与设备
# --------------------------------------------------------------------------


def test_personnel_and_equipment_tables_are_rendered(tmp_path: Path, archive: Path) -> None:
    ctx = context(
        personnel=[
            {"full_name": "张三", "role_code": "lead_inspector", "organization": "某某院",
             "professional_title": "高工", "qualification_certificate_no": "JC-001"}
        ],
        equipment=[
            {"equipment_name": "裂缝观测仪", "model_spec": "ZBL-F130",
             "asset_number": "SB-01", "measurement_range": "0-6mm", "accuracy": "0.01mm",
             "calibration_certificate_no": "JD-2026-1",
             "calibration_valid_until": "2027-01-31", "purpose": "裂缝宽度"}
        ],
    )
    result, _ = build(tmp_path, archive, ctx, anchors=("PERSONNEL_TABLE", "EQUIPMENT_LIST"))

    personnel, equipment = tables(result.output_path)
    assert personnel.rows[1].cells[1].text == "张三"
    assert personnel.rows[1].cells[5].text == "检测负责人"
    assert equipment.rows[1].cells[1].text == "裂缝观测仪"


# --------------------------------------------------------------------------
# 历史对比、评定与结论（设计 §12-§14）
# --------------------------------------------------------------------------


def test_comparison_states_counts_and_carries_the_caveat(
    tmp_path: Path, archive: Path
) -> None:
    """条数变化不能证明病害身份关系，免责说明必须跟着结论走（设计 §12.2）。"""
    ctx = context(
        parts=[
            part(
                "SUPERSTRUCTURE", "上部结构", rows=[defect_row(1)], photos=[],
                comparison=comparison(current=23, previous=18),
            )
        ]
    )
    result, _ = build(
        tmp_path, archive, ctx, anchors=("PREVIOUS_COMPARISON:SUPERSTRUCTURE",)
    )

    text = body_text(result.output_path)
    assert any("2025 年上部结构共记录病害 18 条" in line for line in text)
    assert any("增加 5 条" in line for line in text)
    assert COMPARISON_CAVEAT in text


def test_comparison_never_claims_new_or_repaired_defects(
    tmp_path: Path, archive: Path
) -> None:
    """设计 §12.2 点名禁止的四种说法，一个都不许出现在结论句里。

    免责说明本身要把这几种说法点出来才能否掉它们，所以只查结论句。
    """
    ctx = context(
        parts=[
            part(
                "SUPERSTRUCTURE", "上部结构", rows=[defect_row(1)], photos=[],
                comparison=comparison(current=23, previous=18),
            )
        ]
    )
    result, _ = build(
        tmp_path, archive, ctx, anchors=("PREVIOUS_COMPARISON:SUPERSTRUCTURE",)
    )

    claims = "".join(
        line for line in body_text(result.output_path) if line != COMPARISON_CAVEAT
    )
    for forbidden in ("新增", "修复", "消失", "病害发展"):
        assert forbidden not in claims


def test_comparison_without_history_says_so(tmp_path: Path, archive: Path) -> None:
    result, _ = build(
        tmp_path,
        archive,
        superstructure([defect_row(1)], []),
        anchors=("PREVIOUS_COMPARISON:SUPERSTRUCTURE",),
    )

    assert NO_COMPARISON_TEXT in body_text(result.output_path)


def test_weight_table_lists_absent_components_too(tmp_path: Path, archive: Path) -> None:
    """本桥没有的部件也要列出来并注明"无此构件"（表4.1-1）。

    那一行的权重正是被摊给同部位其余部件的那部分；不列出来，读者看不懂重分配后的
    数字是怎么来的。
    """
    result, _ = build(tmp_path, archive, context(), anchors=("COMPONENT_WEIGHTS",))

    table = tables(result.output_path)[0]
    assert [cell.text for cell in table.rows[0].cells] == list(WEIGHT_TABLE_HEADERS)
    absent = [row for row in table.rows[1:] if row.cells[2].text == "调治构造物"][0]
    assert absent.cells[3].text == "0.02"  # 规范原表权重照印
    assert absent.cells[4].text == "/"  # 没有重新分配后权重
    assert absent.cells[6].text == "无此构件"
    # 有构件的行照常给出重分配后权重和数量。
    present = [row for row in table.rows[1:] if row.cells[2].text == "桥墩"][0]
    assert present.cells[4].text == "1.00"
    assert present.cells[5].text == "3"


def test_weight_column_prints_what_the_engine_used_without_adjustment(
    tmp_path: Path, archive: Path
) -> None:
    """权重列是评定算分时用的那个值四舍五入到两位，不为了凑合计做任何调整。

    参考报告把 0.2857 手工压成 0.28 以让那一列正好合计 1.00；本系统不这么做——
    报告写的必须是程序算的。代价是这一列可能合计 1.01（下面就是这种情形），
    这不是错，是全精度值取两位小数的必然结果。

    这条测试锁的是这个决定：以后改评定算分逻辑，报告层不用跟着改；反过来，
    谁想在报告层"修正"这一列，会先在这里撞墙。
    """
    # 下部结构：调治构造物（0.02）无此构件，其余六项按 0.98 归一，与引擎同一公式。
    configured = [0.02, 0.01, 0.30, 0.30, 0.28, 0.07]
    rows = [
        {"part_code": "SUBSTRUCTURE", "part_label": "下部结构", "order": order,
         "category_id": f"XB{order:02d}", "category_name": f"部件{order}",
         "configured_weight": value, "effective_weight": value / 0.98,
         "component_count": 1, "present": True}
        for order, value in enumerate(configured, start=1)
    ]
    rows.append(
        {"part_code": "SUBSTRUCTURE", "part_label": "下部结构", "order": 7,
         "category_id": "XB07", "category_name": "调治构造物",
         "configured_weight": 0.02, "effective_weight": None,
         "component_count": None, "present": False}
    )
    ctx = context(assessment=assessment(component_weights=rows))

    result, _ = build(tmp_path, archive, ctx, anchors=("COMPONENT_WEIGHTS",))

    table = tables(result.output_path)[0]
    printed = [row.cells[4].text for row in table.rows[1:]]
    assert printed == ["0.02", "0.01", "0.31", "0.31", "0.29", "0.07", "/"]
    # 合计 1.01：全精度值取两位小数的结果，没有被调平。
    assert sum(float(value) for value in printed[:-1]) == pytest.approx(1.01)


def test_weight_table_merges_the_structure_column(tmp_path: Path, archive: Path) -> None:
    result, _ = build(tmp_path, archive, context(), anchors=("COMPONENT_WEIGHTS",))

    table = tables(result.output_path)[0]
    # 上部结构占 2 行合并成一格：两行读到的是同一个单元格。
    assert table.cell(1, 0).text == "上部结构"
    assert table.cell(1, 0)._tc is table.cell(2, 0)._tc
    assert table.cell(3, 0).text == "下部结构"


def test_assessment_table_follows_the_official_layout(
    tmp_path: Path, archive: Path
) -> None:
    """表4.1-2：每个评价部件下按构件评分分档，一档一行。"""
    result, _ = build(tmp_path, archive, context(), anchors=("ASSESSMENT_RESULT",))

    table = tables(result.output_path)[0]
    assert [cell.text for cell in table.rows[0].cells] == list(RESULT_TABLE_HEADERS)
    # 上部承重构件三档：2 个 65 分、4 个 75 分、6 个 100 分。
    assert [row.cells[3].text for row in table.rows[1:4]] == ["2", "4", "6"]
    assert [row.cells[4].text for row in table.rows[1:4]] == ["65.0", "75.0", "100.0"]
    # 部件评分跨这三行合并成一格。
    assert table.cell(1, 5).text == "85.5"
    assert table.cell(1, 5)._tc is table.cell(3, 5)._tc


def test_assessment_table_merges_structure_and_overall_columns(
    tmp_path: Path, archive: Path
) -> None:
    result, _ = build(tmp_path, archive, context(), anchors=("ASSESSMENT_RESULT",))

    table = tables(result.output_path)[0]
    last = len(table.rows) - 1
    # 结构评分、权重、等级按结构合并；上部结构占前 5 行（3 档 + 2 档）。
    assert table.cell(1, 6).text == "85.5"
    assert table.cell(1, 6)._tc is table.cell(5, 6)._tc
    assert table.cell(1, 7).text == "0.40"
    assert table.cell(1, 8).text == "2"  # 等级列只印数字
    # 全桥评分与综合评级贯通全表。
    assert table.cell(1, 9).text == "86.4"
    assert table.cell(1, 9)._tc is table.cell(last, 9)._tc
    assert table.cell(1, 10).text == "2类"  # 综合评级带"类"


def test_assessment_table_is_followed_by_the_summary_sentence(
    tmp_path: Path, archive: Path
) -> None:
    result, _ = build(tmp_path, archive, context(), anchors=("ASSESSMENT_RESULT",))

    assert any(
        "由以上评定过程可知，该桥评定为 2类，处于“有轻微缺损，对桥梁使用功能无影响”。" == line
        for line in body_text(result.output_path)
    )


def test_table_numbers_share_a_sequence_when_the_format_matches(
    tmp_path: Path, archive: Path
) -> None:
    """4.1.1 与 4.1.2 都配「表4.1-{n}」，于是得到 表4.1-1 和 表4.1-2（设计 §7.5）。"""
    result, _ = build(
        tmp_path, archive, context(), anchors=("COMPONENT_WEIGHTS", "ASSESSMENT_RESULT")
    )

    text = body_text(result.output_path)
    assert any(line.startswith("表4.1-1  桥梁部件权重计算表") for line in text)
    assert any(line.startswith("表4.1-2  总体技术状况评定表") for line in text)


def test_control_indicator_says_none_applies_when_none_triggered(
    tmp_path: Path, archive: Path
) -> None:
    result, _ = build(tmp_path, archive, context(), anchors=("CONTROL_INDICATOR",))

    assert NO_CONTROL_INDICATOR_TEXT in body_text(result.output_path)


def test_control_indicator_lists_what_the_engine_recorded(
    tmp_path: Path, archive: Path
) -> None:
    """触发了就把评定引擎记下来的原样列出，不改写、不归纳。"""
    ctx = context(
        assessment=assessment(
            triggered_controls=[
                {"rule_id": "h21.control.main_component",
                 "message": "上部主要构件评分低于 40 分。", "grade_after": "5类"}
            ]
        )
    )
    result, _ = build(tmp_path, archive, ctx, anchors=("CONTROL_INDICATOR",))

    text = body_text(result.output_path)
    assert any("上部主要构件评分低于 40 分。（评定为 5类）" in line for line in text)
    assert NO_CONTROL_INDICATOR_TEXT not in text


def test_overall_assessment_states_score_grade_and_description(
    tmp_path: Path, archive: Path
) -> None:
    result, _ = build(tmp_path, archive, context(), anchors=("OVERALL_ASSESSMENT",))

    assert any(
        "综合该桥技术状况评分及单项控制指标，百股大桥总体技术状况评分为 86.4 分，"
        "评定为 2类，处于“有轻微缺损，对桥梁使用功能无影响”。" == line
        for line in body_text(result.output_path)
    )


def test_appendix_is_a_card_not_a_summary_table(tmp_path: Path, archive: Path) -> None:
    """附表1 是固定版式的评定卡片：上半基本信息、中间十六个部件的等级、下半签署栏。

    不是按部件汇总的数据表——那种表 4.1.2 已经出过一次了。
    """
    result, _ = build(tmp_path, archive, context(), anchors=("ASSESSMENT_APPENDIX",))

    table = tables(result.output_path)[0]
    text = [[cell.text for cell in row.cells] for row in table.rows]
    flat = ["".join(row) for row in text]

    assert table.rows[0].cells[0].text == "桥梁编码"
    assert any("桥梁名称" in line and "百股大桥" in line for line in flat)
    # 中间那段的表头。
    assert any("桥梁组成及评级" in line and "桥梁部件及评级" in line for line in flat)
    # 下半的固定栏目。
    assert any("桥梁总体技术状况评分 Dr" in line and "86.4" in line for line in flat)
    assert any("综合考虑主要部件最差损坏状况最终评定桥梁技术等级" in line for line in flat)
    assert any("记录人" in line and "负责人" in line for line in flat)


def test_appendix_card_lists_every_standard_component_with_its_grade(
    tmp_path: Path, archive: Path
) -> None:
    """十六个规范部件逐项列出，本桥没有的写斜杠——与 表4.1-1 的「无此构件」同一件事。"""
    result, _ = build(tmp_path, archive, context(), anchors=("ASSESSMENT_APPENDIX",))

    table = tables(result.output_path)[0]
    rows = [[cell.text for cell in row.cells] for row in table.rows]
    named = {
        row[5]: row[10]
        for row in rows
        if row[0].isdigit()
    }
    # 等级只印数字，不带"类"。
    assert named["上部承重构件"] == "2"
    assert named["桥墩"] == "1"
    # 调治构造物本桥没有。
    assert named["调治构造物"] == "/"


def test_appendix_card_leaves_missing_archive_fields_blank(
    tmp_path: Path, archive: Path
) -> None:
    """档案里没有的项留空，不编数据（设计 §14 第 5 条）。"""
    ctx = context(bridge_profile={"station_mark": "K109+747"})
    result, _ = build(tmp_path, archive, ctx, anchors=("ASSESSMENT_APPENDIX",))

    rows = [[cell.text for cell in row.cells] for row in tables(result.output_path)[0].rows]
    by_label = {}
    for row in rows[:4]:
        for pair in range(3):
            by_label[row[pair * 4]] = row[pair * 4 + 2]

    assert by_label["桥位桩号"] == "K109+747"
    assert by_label["主要结构"] == ""
    assert by_label["管养单位"] == ""
    assert by_label["上次大中修日期"] == ""


def card_rows(table) -> list[list[str]]:
    """把每行折成不重复的单元格文字。

    合并单元格在 row.cells 里按它跨的每一列各出现一次，直接读会拿到一串重复；
    按底层元素去重才看得出这行真正有几格。
    """
    rows = []
    for row in table.rows:
        texts: list[str] = []
        previous = None
        for cell in row.cells:
            if cell._tc is previous:
                continue
            previous = cell._tc
            texts.append(cell.text)
        rows.append(texts)
    return rows


def card_fields(table) -> dict[str, str]:
    """卡片上的「名称 -> 取值」。每个字段都是「编号、名称、取值」三格挨着。"""
    fields: dict[str, str] = {}
    for row in card_rows(table):
        for index, text in enumerate(row):
            if text.isdigit() and index + 2 < len(row):
                fields[row[index + 1]] = row[index + 2]
    return fields


def test_bridge_card_prints_all_nine_sections_in_order(tmp_path: Path, archive: Path) -> None:
    """附录2 是 A 到 I 九段的桥梁基本状况卡片，段序照正式报告。"""
    result, _ = build(tmp_path, archive, context(), anchors=("BRIDGE_CARD",))

    heads = [
        row[0][:1]
        for row in card_rows(tables(result.output_path)[0])
        if len(row) == 1 and row[0][:1].isalpha()
    ]
    assert heads == ["A", "B", "C", "D", "D", "E", "F", "G", "H", "I"]


def test_bridge_card_numbers_every_one_of_its_ninety_six_cells(
    tmp_path: Path, archive: Path
) -> None:
    """九十六格一格不少。漏一格版式看不出来，卡片却对不上正式报告的编号。"""
    table = tables(
        build(tmp_path, archive, context(), anchors=("BRIDGE_CARD",))[0].output_path
    )[0]
    printed = {text for row in card_rows(table) for text in row}

    assert {str(number) for number in range(1, 97)} <= printed


def test_bridge_card_fills_the_cells_the_archive_can_answer(
    tmp_path: Path, archive: Path
) -> None:
    result, _ = build(tmp_path, archive, context(), anchors=("BRIDGE_CARD",))

    fields = card_fields(tables(result.output_path)[0])
    assert fields["路线编号"] == "S320"
    assert fields["路线名称"] == "大养线"
    assert fields["桥梁编号"] == "L0123"
    assert fields["桥梁名称"] == "百股大桥"
    assert fields["桥位桩号"] == "K12+345"
    assert fields["建成年限"] == "1998"
    assert fields["桥梁全长(m)"] == "68.5"
    assert fields["桥梁分孔（m）"] == "5×13"
    assert fields["结构体系"] == "钢筋混凝土简支板桥"
    # 档案里只有一个管养单位，落在「养护单位」上。
    assert fields["养护单位"] == "太和公路段"


def test_bridge_card_leaves_cells_without_a_data_source_blank(
    tmp_path: Path, archive: Path
) -> None:
    """九十六格里大多数库里还没有对应字段，留空等档案补，不编数据（设计 §14 第 5 条）。"""
    result, _ = build(tmp_path, archive, context(), anchors=("BRIDGE_CARD",))

    fields = card_fields(tables(result.output_path)[0])
    for label in ("路线等级", "设计荷载", "桥面铺装", "设计图纸", "主梁", "填卡人"):
        assert fields[label] == "", label


def test_bridge_card_records_this_inspection_as_the_first_history_row(
    tmp_path: Path, archive: Path
) -> None:
    """F 段第一行由本次检查填，其余行留给手工补历次记录——库里凑不出完整评定史。"""
    result, _ = build(tmp_path, archive, context(), anchors=("BRIDGE_CARD",))

    rows = card_rows(tables(result.output_path)[0])
    start = rows.index(["75", "76", "77", "78", "79"])
    assert rows[start + 2] == ["2026-05-18", "定期检查", "2类", "", ""]
    assert rows[start + 3] == ["", "", "", "", ""]


def test_appendix_cards_use_only_the_small_card_styles(
    tmp_path: Path, archive: Path
) -> None:
    """两张附录卡片用宋体小五的卡片样式，不用正文表格的五号样式（格子排不下）。"""
    result, _ = build(
        tmp_path, archive, context(), anchors=("ASSESSMENT_APPENDIX", "BRIDGE_CARD")
    )

    for table in tables(result.output_path):
        used = {
            paragraph.style.name
            for row in table.rows
            for cell in row.cells
            for paragraph in cell.paragraphs
        }
        assert used <= {STYLE_CARD_BAND, STYLE_CARD_CELL, STYLE_CARD_HEADER}


def test_conclusion_states_grade_parts_defects_and_advice(
    tmp_path: Path, archive: Path
) -> None:
    result, _ = build(tmp_path, archive, context(), anchors=("CONCLUSION",))

    text = body_text(result.output_path)
    assert any("全桥技术状况评分 86.4 分，评定为 2类" in line for line in text)
    assert any("上部结构 85.5 分（2类）" in line for line in text)
    assert any("1#跨桥面铺装坑槽（扣 12.0 分）" in line for line in text)
    assert MAINTENANCE_ADVICE["2类"] in text


def test_conclusion_never_invents_cost_or_safety_claims(
    tmp_path: Path, archive: Path
) -> None:
    """设计 §14 第 5 条：工程量、预算、工期、结构安全结论都没有数据依据。"""
    result, _ = build(tmp_path, archive, context(), anchors=("CONCLUSION",))

    joined = "".join(body_text(result.output_path))
    for forbidden in ("工程量", "预算", "万元", "工期", "承载能力", "安全性"):
        assert forbidden not in joined


def test_blocks_needing_assessment_abort_without_a_formal_run(
    tmp_path: Path, archive: Path
) -> None:
    """评定在生成期间被撤了，绝不出一份没有评分的定期检测报告。"""
    ctx = context(assessment=dict(NO_ASSESSMENT))

    for block in ("ASSESSMENT_RESULT", "ASSESSMENT_APPENDIX", "CONCLUSION"):
        with pytest.raises(ReportBuildError) as raised:
            build(tmp_path, archive, ctx, anchors=(block,))
        assert raised.value.code == "report_assessment_missing"


def test_bridge_profile_is_prose_that_shortens_without_data(
    tmp_path: Path, archive: Path
) -> None:
    """§1.1 出的是几段话，不是两列表（两列表在附录2 的卡片里）。

    档案缺的项直接不出现，既不留「未知」，也不留只剩标签的半截句子。
    """
    ctx = context(bridge_profile={"bridge_scale": "中桥", "built_year": 1998})
    result, _ = build(tmp_path, archive, ctx, anchors=("BRIDGE_PROFILE",))

    assert tables(result.output_path) == []
    text = _all_text(result.output_path)
    assert "百股大桥位于大养线公路（S320）太和区段，建成于 1998 年。属中桥。" in text
    assert "未知" not in text
    assert "设计单位" not in text


# --------------------------------------------------------------------------
# 装配纪律
# --------------------------------------------------------------------------


def test_no_anchor_survives_the_build(tmp_path: Path, archive: Path) -> None:
    photos = [photo("照片2.1-1", ARCHIVE)]
    ctx = superstructure([defect_row(1, photo_numbers=["照片2.1-1"])], photos)
    result, _ = build(tmp_path, archive, ctx)

    assert "[[REPORT:" not in _all_text(result.output_path)
    assert result.blocks_rendered == list(
        (
            "PERSONNEL_TABLE",
            "EQUIPMENT_LIST",
            "DEFECT_TABLES:SUPERSTRUCTURE",
            "DEFECT_PHOTOS:SUPERSTRUCTURE",
            "DEFECT_TABLES:SUBSTRUCTURE",
            "DEFECT_PHOTOS:SUBSTRUCTURE",
        )
    )


def test_every_contract_block_has_a_renderer() -> None:
    """契约要求的每个内容块都装得出来，没有留白的缺口。

    契约里加了新内容块而没写装配规则时，这条会先炸——比让它在生成一半时报
    report_block_not_implemented 早得多。
    """
    assert PERIODIC_INSPECTION_V1.all_anchors <= set(RENDERERS)
    assert PENDING_BLOCKS == frozenset()


def test_unknown_block_aborts_instead_of_leaving_a_hole(
    tmp_path: Path, archive: Path
) -> None:
    """装不出来的块必须报错，不能悄悄留个空段落。"""
    with pytest.raises(ReportBuildError) as raised:
        build(tmp_path, archive, superstructure([], []), anchors=("NOT_A_BLOCK",))

    assert raised.value.code == "report_block_unknown"


def test_template_file_is_not_modified(tmp_path: Path, archive: Path) -> None:
    photos = [photo("照片2.1-1", ARCHIVE)]
    ctx = superstructure([defect_row(1, photo_numbers=["照片2.1-1"])], photos)
    template = build_builder_template(tmp_path / "template.docx")
    before = template.read_bytes()

    build_report(template, ctx, tmp_path / "out.docx", archive)

    assert template.read_bytes() == before


def test_building_over_the_template_is_refused(tmp_path: Path, archive: Path) -> None:
    template = build_builder_template(tmp_path / "template.docx")

    with pytest.raises(ReportBuildError) as raised:
        build_report(template, superstructure([], []), template, archive)

    assert raised.value.code == "report_output_overwrites_template"


def test_build_is_deterministic(tmp_path: Path, archive: Path) -> None:
    """同一份上下文生成两次，文档文字必须一致（设计 §5.5）。"""
    photos = [photo(f"照片2.1-{n}", ARCHIVE) for n in range(1, 3)]
    ctx = superstructure(
        [defect_row(1, photo_numbers=[p["report_number"] for p in photos])], photos
    )
    template = build_builder_template(tmp_path / "template.docx")

    first = build_report(template, ctx, tmp_path / "a.docx", archive)
    second = build_report(template, ctx, tmp_path / "b.docx", archive)

    assert _all_text(first.output_path) == _all_text(second.output_path)


# --------------------------------------------------------------------------


def _all_text(path: Path) -> str:
    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _picture_count(path: Path) -> int:
    return len(Document(str(path)).element.body.findall(f".//{qn('pic:pic')}"))


def _first_picture_extent(path: Path) -> tuple[int, int] | None:
    document = Document(str(path))
    for extent in document.element.body.iter(qn("wp:extent")):
        return int(extent.get("cx")), int(extent.get("cy"))
    return None
