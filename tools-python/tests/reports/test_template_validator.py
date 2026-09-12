from __future__ import annotations

from pathlib import Path

from docx import Document

from bridge_report_tools.reports.template_validator import (
    TemplateConfig,
    validate_template,
)
from tests.reports.template_fixtures import (
    CORE_ANCHORS,
    NUMBER_FORMATS,
    add_complex_field,
    add_simple_field,
    add_split_text,
    build_template,
    valid_config,
)


def test_minimal_template_with_full_config_is_valid(tmp_path: Path) -> None:
    result = validate_template(build_template(tmp_path / "t.docx"), valid_config())

    assert result.is_valid, result.codes()
    assert result.anchors_in_document_order == list(CORE_ANCHORS)
    assert result.placeholders_used == ["bridge_name", "report_no"]
    assert result.fields_used == ["PAGEREF", "TOC"]


def test_template_order_is_reported_as_written_not_sorted(tmp_path: Path) -> None:
    """模板决定顺序：校验结果必须反映文档里的真实次序（设计 §7.4）。"""
    reordered = (
        "ASSESSMENT_RESULT",
        "DEFECT_PHOTOS:DECK",
        "DEFECT_TABLES:DECK",
        "PREVIOUS_COMPARISON:DECK",
        "DEFECT_PHOTOS:SUPERSTRUCTURE",
        "DEFECT_TABLES:SUPERSTRUCTURE",
        "PREVIOUS_COMPARISON:SUPERSTRUCTURE",
        "DEFECT_PHOTOS:SUBSTRUCTURE",
        "DEFECT_TABLES:SUBSTRUCTURE",
        "PREVIOUS_COMPARISON:SUBSTRUCTURE",
        "BRIDGE_PROFILE",
        "COMPONENT_WEIGHTS",
        "CONTROL_INDICATOR",
        "OVERALL_ASSESSMENT",
        "CONCLUSION",
        "ASSESSMENT_APPENDIX",
        "BRIDGE_CARD",
    )
    result = validate_template(
        build_template(tmp_path / "t.docx", anchors=reordered), valid_config()
    )

    assert result.is_valid, result.codes()
    assert result.anchors_in_document_order == list(reordered)


def test_missing_core_anchor_is_invalid(tmp_path: Path) -> None:
    without_conclusion = [a for a in CORE_ANCHORS if a != "CONCLUSION"]
    result = validate_template(
        build_template(tmp_path / "t.docx", anchors=without_conclusion), valid_config()
    )

    assert not result.is_valid
    assert "template_anchor_missing" in result.codes()


def test_duplicate_anchor_is_invalid(tmp_path: Path) -> None:
    result = validate_template(
        build_template(tmp_path / "t.docx", anchors=(*CORE_ANCHORS, "DEFECT_TABLES:SUPERSTRUCTURE")),
        valid_config(),
    )

    assert not result.is_valid
    assert "template_anchor_duplicated" in result.codes()


def test_unknown_anchor_is_invalid(tmp_path: Path) -> None:
    result = validate_template(
        build_template(tmp_path / "t.docx", anchors=(*CORE_ANCHORS, "WEATHER_SUMMARY")),
        valid_config(),
    )

    assert not result.is_valid
    assert "template_anchor_unknown" in result.codes()


def test_anchor_sharing_paragraph_is_invalid(tmp_path: Path) -> None:
    def customize(document: Document) -> None:
        document.add_paragraph("见下表 [[REPORT:PERSONNEL_TABLE]] 完")

    result = validate_template(
        build_template(tmp_path / "t.docx", customize=customize), valid_config()
    )

    assert not result.is_valid
    assert "template_anchor_not_own_paragraph" in result.codes()


def test_anchor_in_header_is_invalid(tmp_path: Path) -> None:
    def customize(document: Document) -> None:
        header = document.sections[0].header
        header.is_linked_to_previous = False
        header.paragraphs[0].text = "[[REPORT:EQUIPMENT_LIST]]"

    result = validate_template(
        build_template(tmp_path / "t.docx", customize=customize), valid_config()
    )

    assert not result.is_valid
    assert "template_anchor_outside_body" in result.codes()


def test_unknown_placeholder_is_invalid(tmp_path: Path) -> None:
    result = validate_template(
        build_template(tmp_path / "t.docx", placeholders=("report_no", "weather")),
        valid_config(),
    )

    assert not result.is_valid
    assert "template_placeholder_unknown" in result.codes()


def test_split_placeholder_is_recognised_not_reported(tmp_path: Path) -> None:
    """跨 run 的占位符必须被认出来，否则会被误判为未知占位符。"""

    def customize(document: Document) -> None:
        add_split_text(document.add_paragraph(), "{{admini", "strative_", "region}}")

    result = validate_template(
        build_template(tmp_path / "t.docx", customize=customize), valid_config()
    )

    assert result.is_valid, result.codes()
    assert "administrative_region" in result.placeholders_used


def test_unbalanced_braces_are_invalid(tmp_path: Path) -> None:
    def customize(document: Document) -> None:
        document.add_paragraph("{{report_no 漏了右括号")

    result = validate_template(
        build_template(tmp_path / "t.docx", customize=customize), valid_config()
    )

    assert not result.is_valid
    assert "template_placeholder_unbalanced" in result.codes()


def test_seq_field_is_rejected(tmp_path: Path) -> None:
    def customize(document: Document) -> None:
        add_complex_field(document.add_paragraph(), " SEQ 表 \\* ARABIC ", "1")

    result = validate_template(
        build_template(tmp_path / "t.docx", customize=customize), valid_config()
    )

    assert not result.is_valid
    assert "template_field_rejected" in result.codes()


def test_styleref_field_is_rejected(tmp_path: Path) -> None:
    def customize(document: Document) -> None:
        header = document.sections[0].header
        header.is_linked_to_previous = False
        add_complex_field(header.paragraphs[0], ' STYLEREF "标题 1" \\* MERGEFORMAT ', "第一章")

    result = validate_template(
        build_template(tmp_path / "t.docx", customize=customize), valid_config()
    )

    assert not result.is_valid
    assert "template_field_rejected" in result.codes()


def test_standalone_pageref_is_rejected_but_toc_internal_one_is_not(tmp_path: Path) -> None:
    def customize(document: Document) -> None:
        add_complex_field(document.add_paragraph(), " PAGEREF _Ref9 \\h ", "12")

    result = validate_template(
        build_template(tmp_path / "t.docx", customize=customize), valid_config()
    )

    assert not result.is_valid
    # 同一份文档里 TOC 内部也有一个 PAGEREF，它不能被一起判错。
    assert result.codes().count("template_field_outside_toc") == 1


def test_page_and_numpages_fields_are_allowed(tmp_path: Path) -> None:
    def customize(document: Document) -> None:
        footer = document.sections[0].footer
        footer.is_linked_to_previous = False
        add_simple_field(footer.paragraphs[0], " PAGE ", "1")
        add_simple_field(footer.paragraphs[0], " NUMPAGES ", "34")

    result = validate_template(
        build_template(tmp_path / "t.docx", customize=customize), valid_config()
    )

    assert result.is_valid, result.codes()


def test_unlisted_field_is_rejected(tmp_path: Path) -> None:
    def customize(document: Document) -> None:
        add_simple_field(document.add_paragraph(), " DATE \\@ &quot;yyyy&quot; ", "2026")

    result = validate_template(
        build_template(tmp_path / "t.docx", customize=customize), valid_config()
    )

    assert not result.is_valid
    assert "template_field_not_allowed" in result.codes()


def test_photo_block_requires_a_number_format(tmp_path: Path) -> None:
    """报告图号按模板结构现编，所以每个照片块都必须配格式（设计 §7.5）。

    库里的 defect_photos.photo_number 不是可用的图号：它由导入器按病害 UUID 顺序
    现编，与构件台账顺序对不上。模板不配格式就无号可用，必须在这里拦住。
    """
    config = valid_config(
        table_number_formats={
            key: value
            for key, value in NUMBER_FORMATS.items()
            if not key.startswith("DEFECT_PHOTOS")
        }
    )

    result = validate_template(build_template(tmp_path / "t.docx"), config)

    assert not result.is_valid
    assert result.codes().count("template_number_format_missing") == 3


def test_missing_number_format_is_invalid(tmp_path: Path) -> None:
    config = valid_config(
        table_number_formats={"DEFECT_TABLES:SUPERSTRUCTURE": "表2.1-{n}"}
    )

    result = validate_template(build_template(tmp_path / "t.docx"), config)

    assert not result.is_valid
    assert result.codes().count("template_number_format_missing") == 8


def test_number_format_without_token_is_invalid(tmp_path: Path) -> None:
    config = valid_config(
        table_number_formats={**NUMBER_FORMATS, "DEFECT_TABLES:DECK": "表2.3-1"}
    )

    result = validate_template(build_template(tmp_path / "t.docx"), config)

    assert not result.is_valid
    assert "template_number_format_invalid" in result.codes()


def test_number_format_with_two_tokens_is_invalid(tmp_path: Path) -> None:
    config = valid_config(
        table_number_formats={**NUMBER_FORMATS, "DEFECT_TABLES:DECK": "表{n}.3-{n}"}
    )

    result = validate_template(build_template(tmp_path / "t.docx"), config)

    assert not result.is_valid
    assert "template_number_format_invalid" in result.codes()


def test_unknown_personnel_role_is_invalid(tmp_path: Path) -> None:
    config = valid_config(required_personnel_roles=["approver", "chief_engineer"])

    result = validate_template(build_template(tmp_path / "t.docx"), config)

    assert not result.is_valid
    assert "template_personnel_role_unknown" in result.codes()


def test_personnel_placeholder_needs_declared_role(tmp_path: Path) -> None:
    """声明缺失时生成前检查放行、签字页却留空，属于静默错误输出。"""
    config = valid_config(required_personnel_roles=["approver"])

    result = validate_template(
        build_template(
            tmp_path / "t.docx",
            placeholders=("report_no", "personnel.reviewer.names"),
        ),
        config,
    )

    assert not result.is_valid
    assert "template_personnel_role_not_required" in result.codes()


def test_declared_personnel_placeholder_is_valid(tmp_path: Path) -> None:
    result = validate_template(
        build_template(
            tmp_path / "t.docx",
            placeholders=("report_no", "personnel.approver.names"),
        ),
        valid_config(),
    )

    assert result.is_valid, result.codes()


def test_default_config_reports_every_missing_number_format(tmp_path: Path) -> None:
    result = validate_template(build_template(tmp_path / "t.docx"), TemplateConfig())

    assert not result.is_valid
    assert result.codes().count("template_number_format_missing") == 9


def test_per_part_block_without_part_is_invalid(tmp_path: Path) -> None:
    """病害表在正式报告里是每个结构部位一张（表2.1-1 / 2.2-1 / 2.3-1），不能只写一个。"""
    anchors = [a for a in CORE_ANCHORS if not a.startswith("DEFECT_TABLES")]
    result = validate_template(
        build_template(tmp_path / "t.docx", anchors=(*anchors, "DEFECT_TABLES")),
        valid_config(),
    )

    assert not result.is_valid
    assert "template_anchor_part_required" in result.codes()


def test_unknown_structure_part_code_is_invalid(tmp_path: Path) -> None:
    result = validate_template(
        build_template(tmp_path / "t.docx", anchors=(*CORE_ANCHORS, "DEFECT_TABLES:PIER")),
        valid_config(),
    )

    assert not result.is_valid
    assert "template_anchor_part_unknown" in result.codes()


def test_part_suffix_on_non_per_part_block_is_invalid(tmp_path: Path) -> None:
    anchors = [a for a in CORE_ANCHORS if a != "CONCLUSION"]
    result = validate_template(
        build_template(tmp_path / "t.docx", anchors=(*anchors, "CONCLUSION:DECK")),
        valid_config(),
    )

    assert not result.is_valid
    assert "template_anchor_part_not_allowed" in result.codes()


def test_inconsistent_part_coverage_is_invalid(tmp_path: Path) -> None:
    """病害表分了三部位、照片只放了上部——桥面系的照片会无处可去。"""
    anchors = [a for a in CORE_ANCHORS if not a.startswith("DEFECT_PHOTOS")]
    result = validate_template(
        build_template(
            tmp_path / "t.docx", anchors=(*anchors, "DEFECT_PHOTOS:SUPERSTRUCTURE")
        ),
        valid_config(),
    )

    assert not result.is_valid
    assert "template_part_coverage_mismatch" in result.codes()


def test_consistent_narrower_part_coverage_is_valid(tmp_path: Path) -> None:
    """三类内容块一致地只覆盖上部结构时，模板本身是自洽的。

    数据落在模板没有覆盖的部位，属于生成前检查该拦的事，不是模板校验能判断的。
    """
    anchors = [
        "BRIDGE_PROFILE",
        "DEFECT_TABLES:SUPERSTRUCTURE",
        "DEFECT_PHOTOS:SUPERSTRUCTURE",
        "PREVIOUS_COMPARISON:SUPERSTRUCTURE",
        "COMPONENT_WEIGHTS",
        "ASSESSMENT_RESULT",
        "CONTROL_INDICATOR",
        "OVERALL_ASSESSMENT",
        "CONCLUSION",
        "ASSESSMENT_APPENDIX",
        "BRIDGE_CARD",
    ]
    config = valid_config(
        table_number_formats={
            "DEFECT_TABLES:SUPERSTRUCTURE": "表2.1-{n}",
            "DEFECT_PHOTOS:SUPERSTRUCTURE": "照片2.1-{n}",
            "COMPONENT_WEIGHTS": "表4.1-{n}",
            "ASSESSMENT_RESULT": "表4.1-{n}",
            "ASSESSMENT_APPENDIX": "附表1-{n}",
        }
    )

    result = validate_template(build_template(tmp_path / "t.docx", anchors=anchors), config)

    assert result.is_valid, result.codes()


def test_number_format_key_without_part_is_invalid(tmp_path: Path) -> None:
    config = valid_config(
        table_number_formats={**NUMBER_FORMATS, "DEFECT_TABLES": "表 2-{n}"}
    )

    result = validate_template(build_template(tmp_path / "t.docx"), config)

    assert not result.is_valid
    assert "template_number_format_part_required" in result.codes()


def test_number_format_key_with_part_on_plain_block_is_invalid(tmp_path: Path) -> None:
    config = valid_config(
        table_number_formats={**NUMBER_FORMATS, "ASSESSMENT_RESULT:DECK": "表4-{n}"}
    )

    result = validate_template(build_template(tmp_path / "t.docx"), config)

    assert not result.is_valid
    assert "template_number_format_part_not_allowed" in result.codes()
