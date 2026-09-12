"""构造模板 docx 的测试固件。

模板里的域必须用原始 XML 拼——python-docx 没有建域的 API，而域约束正是要测的核心
规则之一（设计 §7.6）。
"""

from __future__ import annotations

from pathlib import Path
from typing import Callable, Iterable

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.text.paragraph import Paragraph

from bridge_report_tools.reports.styles import (
    REQUIRED_BUILDER_STYLES,
    STYLE_PHOTO_LAYOUT,
    STYLE_TABLE,
)
from bridge_report_tools.reports.template_validator import TemplateConfig


#: 与两份正式报告一致：第 2 章按上部、下部、桥面系各起一节，每节自带病害表、照片和对比。
TEMPLATE_PARTS = ("SUPERSTRUCTURE", "SUBSTRUCTURE", "DECK")

CORE_ANCHORS = (
    "BRIDGE_PROFILE",
    "DEFECT_TABLES:SUPERSTRUCTURE",
    "DEFECT_PHOTOS:SUPERSTRUCTURE",
    "PREVIOUS_COMPARISON:SUPERSTRUCTURE",
    "DEFECT_TABLES:SUBSTRUCTURE",
    "DEFECT_PHOTOS:SUBSTRUCTURE",
    "PREVIOUS_COMPARISON:SUBSTRUCTURE",
    "DEFECT_TABLES:DECK",
    "DEFECT_PHOTOS:DECK",
    "PREVIOUS_COMPARISON:DECK",
    "COMPONENT_WEIGHTS",
    "ASSESSMENT_RESULT",
    "CONTROL_INDICATOR",
    "OVERALL_ASSESSMENT",
    "CONCLUSION",
    "ASSESSMENT_APPENDIX",
    "BRIDGE_CARD",
)

NUMBER_FORMATS = {
    "DEFECT_TABLES:SUPERSTRUCTURE": "表2.1-{n}",
    "DEFECT_TABLES:SUBSTRUCTURE": "表2.2-{n}",
    "DEFECT_TABLES:DECK": "表2.3-{n}",
    "DEFECT_PHOTOS:SUPERSTRUCTURE": "照片2.1-{n}",
    "DEFECT_PHOTOS:SUBSTRUCTURE": "照片2.2-{n}",
    "DEFECT_PHOTOS:DECK": "照片2.3-{n}",
    # 4.1.1 与 4.1.2 共用一条序列：格式串相同即共号。
    "COMPONENT_WEIGHTS": "表4.1-{n}",
    "ASSESSMENT_RESULT": "表4.1-{n}",
    "ASSESSMENT_APPENDIX": "附表1-{n}",
}


def valid_config(**overrides) -> TemplateConfig:
    payload = {
        "table_number_formats": dict(NUMBER_FORMATS),
        "required_personnel_roles": ["approver", "reviewer", "lead_inspector", "compiler"],
    }
    payload.update(overrides)
    return TemplateConfig(**payload)


def add_split_text(paragraph: Paragraph, *chunks: str) -> Paragraph:
    """把一段文字拆进多个 run，模拟 Word 对占位符的切割。"""
    for chunk in chunks:
        paragraph.add_run(chunk)
    return paragraph


def _append_run_xml(paragraph: Paragraph, inner: str) -> None:
    paragraph._p.append(parse_xml(f'<w:r {nsdecls("w")}>{inner}</w:r>'))


def add_complex_field(
    paragraph: Paragraph,
    instruction: str,
    result_text: str = "",
    nested: list[tuple[str, str]] | None = None,
) -> Paragraph:
    """写一个复杂域；nested 里的域落在本域的结果区内部（用于造 TOC 内部的 PAGEREF）。"""
    _append_run_xml(paragraph, '<w:fldChar w:fldCharType="begin"/>')
    _append_run_xml(
        paragraph,
        f'<w:instrText xml:space="preserve">{instruction}</w:instrText>',
    )
    _append_run_xml(paragraph, '<w:fldChar w:fldCharType="separate"/>')
    for nested_instruction, nested_result in nested or []:
        add_complex_field(paragraph, nested_instruction, nested_result)
    if result_text:
        _append_run_xml(paragraph, f"<w:t>{result_text}</w:t>")
    _append_run_xml(paragraph, '<w:fldChar w:fldCharType="end"/>')
    return paragraph


def add_simple_field(paragraph: Paragraph, instruction: str, result_text: str = "") -> None:
    paragraph._p.append(
        parse_xml(
            f'<w:fldSimple {nsdecls("w")} w:instr="{instruction}">'
            f"<w:r><w:t>{result_text}</w:t></w:r>"
            f"</w:fldSimple>"
        )
    )


def add_toc(paragraph: Paragraph) -> None:
    """一个带内部 PAGEREF 的目录域，形状与 Word 刷新后的产物一致。"""
    add_complex_field(
        paragraph,
        ' TOC \\o "1-3" \\h \\z \\u ',
        result_text="",
        nested=[(" PAGEREF _Toc12345 \\h ", "3")],
    )


def build_template(
    path: Path,
    anchors: Iterable[str] = CORE_ANCHORS,
    placeholders: Iterable[str] = ("report_no", "bridge_name"),
    customize: Callable[[Document], None] | None = None,
) -> Path:
    """按需拼一份模板；默认产出满足契约的最小合规模板。"""
    document = Document()
    for name in placeholders:
        document.add_paragraph(f"{{{{{name}}}}}")
    add_toc(document.add_paragraph())
    for anchor in anchors:
        document.add_paragraph(f"[[REPORT:{anchor}]]")
    if customize is not None:
        customize(document)
    document.save(str(path))
    return path


def minimal_valid_template(path: Path) -> Path:
    return build_template(path)


#: Docx Builder 在第 6 步已实现的内容块（设计 §27 第 6 步）。评定、历史对比和结论
#: 是第 7 步，故意不放进来——放进来就得先造假内容才能跑通。
BUILDER_ANCHORS = (
    "PERSONNEL_TABLE",
    "EQUIPMENT_LIST",
    "DEFECT_TABLES:SUPERSTRUCTURE",
    "DEFECT_PHOTOS:SUPERSTRUCTURE",
    "DEFECT_TABLES:SUBSTRUCTURE",
    "DEFECT_PHOTOS:SUBSTRUCTURE",
)


def build_builder_template(
    path: Path,
    anchors: Iterable[str] = BUILDER_ANCHORS,
    placeholders: Iterable[str] = ("report_no", "bridge_name"),
    customize: Callable[[Document], None] | None = None,
) -> Path:
    """给 Docx Builder 用的模板：带齐生成器要用的命名样式。

    只建样式的名字，不复刻字体字号——Builder 按名字取样式，长什么样是模板作者的事
    （设计 §18）。真正随系统发布那份模板的版式由 test_standard_template 守。
    """
    document = Document()
    for name in sorted(REQUIRED_BUILDER_STYLES):
        style_type = (
            WD_STYLE_TYPE.TABLE
            if name in (STYLE_TABLE, STYLE_PHOTO_LAYOUT)
            else WD_STYLE_TYPE.PARAGRAPH
        )
        document.styles.add_style(name, style_type)
    for name in placeholders:
        document.add_paragraph(f"{{{{{name}}}}}")
    for anchor in anchors:
        document.add_paragraph(f"[[REPORT:{anchor}]]")
    if customize is not None:
        customize(document)
    document.save(str(path))
    return path
