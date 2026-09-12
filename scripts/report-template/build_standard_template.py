"""生成首个标准模板 `periodic_inspection_v1`（设计 §8）。

版式取自 docs/ 下两份正式报告，但样式和域全部重建，不继承旧文档：
两份原件里有 913 个 STYLEREF、458 个 SEQ、455 个 REF，还有把 2.3 节小标题写成
"2.2.3病害成因分析"这类手工编号错误。这些正是本设计要甩掉的历史漂移。

模板只提供骨架：封面、声明页、签字页、目录、章节标题、分节、页眉页脚，以及生成器
写动态内容时要用的命名样式。病害表、照片页、评定表由 Docx Builder 在锚点处装配。

这是一次性的编排工具，不是系统的运行时功能：它只负责把模板画出来。真正随系统发布的
是它的产物——templates/report/periodic_inspection_v1/ 下的 .docx 和 template.json。
模板契约本身（锚点语法、样式名）在 bridge_report_tools.reports 里，这里只是遵守它。

**模板已于 2026-09-10 转为手工维护。** 从那天起 .docx 是唯一真源，版式改动一律在
Word 里做；这个脚本只用来生成初版，目标文件已存在时直接拒绝运行，免得把人改的东西
一次跑没。要重新生成初版，指定一个空目录，再手工比对合并。

用法（在仓库根目录）：
    tools-python/.venv/Scripts/python.exe scripts/report-template/build_standard_template.py <空目录>
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO_ROOT / "tools-python"))

from docx import Document  # noqa: E402
from docx.enum.section import WD_ORIENT, WD_SECTION  # noqa: E402
from docx.enum.style import WD_STYLE_TYPE  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING  # noqa: E402
from docx.oxml import OxmlElement, parse_xml  # noqa: E402
from docx.oxml.ns import nsdecls, qn  # noqa: E402
from docx.shared import Cm, Pt  # noqa: E402

from bridge_report_tools.reports.contract import (  # noqa: E402
    CONTRACT_PERIODIC_INSPECTION_V1,
    anchor_text,
)
from bridge_report_tools.reports.styles import (  # noqa: E402
    STYLE_BODY,
    STYLE_CARD_BAND,
    STYLE_CARD_CELL,
    STYLE_CARD_HEADER,
    STYLE_COVER_FIELD,
    STYLE_COVER_TITLE,
    STYLE_NOTICE,
    STYLE_PHOTO,
    STYLE_PHOTO_CAPTION,
    STYLE_PHOTO_LAYOUT,
    STYLE_TABLE,
    STYLE_TABLE_CAPTION,
    STYLE_TABLE_CELL,
    STYLE_TABLE_HEADER,
)

DEFAULT_OUTPUT_DIR = REPO_ROOT / "templates" / "report" / "periodic_inspection_v1"
TEMPLATE_BASENAME = "periodic-inspection-v1"


# 版面取自百股大桥报告：A4，上下 2.2cm、左右 2.5cm，页眉页脚 1.3cm；附录横向。
PAGE_WIDTH = Cm(21.0)
PAGE_HEIGHT = Cm(29.7)
MARGIN_TOP = Cm(2.2)
MARGIN_BOTTOM = Cm(2.2)
MARGIN_LEFT = Cm(2.5)
MARGIN_RIGHT = Cm(2.5)
HEADER_DISTANCE = Cm(1.3)
FOOTER_DISTANCE = Cm(1.3)
LANDSCAPE_MARGIN = Cm(2.54)

FONT_SONG = "宋体"
FONT_HEI = "黑体"
FONT_KAI = "楷体"
FONT_ASCII = "Times New Roman"

STRUCTURE_SECTIONS = (
    ("2.1", "上部结构", "SUPERSTRUCTURE"),
    ("2.2", "下部结构", "SUBSTRUCTURE"),
    ("2.3", "桥面系", "DECK"),
)

#: 与 STRUCTURE_SECTIONS 对应的表号格式，写进 report_templates.contract_config_json。
TEMPLATE_NUMBER_FORMATS = {
    "DEFECT_TABLES:SUPERSTRUCTURE": "表2.1-{n}",
    "DEFECT_TABLES:SUBSTRUCTURE": "表2.2-{n}",
    "DEFECT_TABLES:DECK": "表2.3-{n}",
    # 报告图号按模板结构现编，与库里的 defect_photos.photo_number 无关（设计 §7.5）。
    "DEFECT_PHOTOS:SUPERSTRUCTURE": "照片2.1-{n}",
    "DEFECT_PHOTOS:SUBSTRUCTURE": "照片2.2-{n}",
    "DEFECT_PHOTOS:DECK": "照片2.3-{n}",
    # 4.1.1 的权重表和 4.1.2 的评定表同属 4.1 节，共用一条编号序列：格式串相同即
    # 共号，于是按文档顺序得到 表4.1-1 和 表4.1-2（设计 §7.5）。
    "COMPONENT_WEIGHTS": "表4.1-{n}",
    "ASSESSMENT_RESULT": "表4.1-{n}",
    "ASSESSMENT_APPENDIX": "附表1-{n}",
}

TEMPLATE_REQUIRED_ROLES = ["approver", "reviewer", "lead_inspector", "compiler"]

NOTICE_CLAUSES = (
    "1、本检测报告加盖检验检测专用章、CMA 章和报告骑缝章后有效，否则无效；",
    "2、本检测报告签字手续不齐全的无效；",
    "3、本检测报告涂改、增删无效；",
    "4、复制的检测报告，未经本单位同意并加盖检验检测专用章的无效；",
    "5、委托单位对检测报告有异议，应当在收到报告之日起十五日内或检测可复现期间向本单位提出，逾期不予受理；",
    "6、对于来样委托试验，仅对来样样品的检测数据负责。",
)


# --------------------------------------------------------------------------
# 低层工具
# --------------------------------------------------------------------------


def _set_style_font(
    style,
    *,
    ascii_font: str = FONT_ASCII,
    east_asia: str = FONT_SONG,
    size_pt: float = 10.5,
    bold: bool = False,
) -> None:
    """样式的中文字体只能落在 w:rPr/w:rFonts 的 w:eastAsia 上，python-docx 不直接支持。"""
    style.font.name = ascii_font
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    fonts.set(qn("w:ascii"), ascii_font)
    fonts.set(qn("w:hAnsi"), ascii_font)
    fonts.set(qn("w:eastAsia"), east_asia)


def _set_first_line_indent(style, characters: int = 2) -> None:
    """按"字符"缩进而不是按厘米：中文正文换字号时缩进要跟着变。"""
    ppr = style.element.get_or_add_pPr()
    indent = ppr.find(qn("w:ind"))
    if indent is None:
        indent = OxmlElement("w:ind")
        ppr.append(indent)
    indent.set(qn("w:firstLineChars"), str(characters * 100))


def _set_page_numbering_restart(section, start: int = 1) -> None:
    sect_pr = section._sectPr
    existing = sect_pr.find(qn("w:pgNumType"))
    if existing is not None:
        sect_pr.remove(existing)
    sect_pr.append(parse_xml(f'<w:pgNumType {nsdecls("w")} w:start="{start}"/>'))


def _add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    """插入一个复杂域。模板只用 TOC / PAGE / NUMPAGES 三种（设计 §7.6）。"""

    def run(inner: str) -> None:
        paragraph._p.append(parse_xml(f'<w:r {nsdecls("w")}>{inner}</w:r>'))

    run('<w:fldChar w:fldCharType="begin"/>')
    run(f'<w:instrText xml:space="preserve">{instruction}</w:instrText>')
    run('<w:fldChar w:fldCharType="separate"/>')
    if placeholder:
        run(f"<w:t>{placeholder}</w:t>")
    run('<w:fldChar w:fldCharType="end"/>')


def _portrait(section) -> None:
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = PAGE_WIDTH
    section.page_height = PAGE_HEIGHT
    section.top_margin = MARGIN_TOP
    section.bottom_margin = MARGIN_BOTTOM
    section.left_margin = MARGIN_LEFT
    section.right_margin = MARGIN_RIGHT
    section.header_distance = HEADER_DISTANCE
    section.footer_distance = FOOTER_DISTANCE


def _landscape(section) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = PAGE_HEIGHT
    section.page_height = PAGE_WIDTH
    section.top_margin = LANDSCAPE_MARGIN
    section.bottom_margin = LANDSCAPE_MARGIN
    section.left_margin = LANDSCAPE_MARGIN
    section.right_margin = LANDSCAPE_MARGIN
    section.header_distance = HEADER_DISTANCE
    section.footer_distance = FOOTER_DISTANCE


def _blank_header_footer(section) -> None:
    """封面、声明页这类不带页眉页脚的节，必须显式断开继承。"""
    for part in (section.header, section.footer):
        part.is_linked_to_previous = False
        for paragraph in part.paragraphs:
            paragraph.text = ""


def _running_header_footer(section) -> None:
    """页眉写静态文字，页脚用 PAGE / NUMPAGES。

    原报告的页眉是 913 个 STYLEREF 域拼出来的动态章节名；这里改成占位符，
    由生成器一次性写死，避免 Word 重排时把编号刷成别的章节。
    """
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.text = "{{project_name}}    {{bridge_name}}定期检测报告    {{administrative_region}}"
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    footer = section.footer
    footer.is_linked_to_previous = False
    foot = footer.paragraphs[0]
    foot.text = ""
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    foot.add_run("第 ")
    _add_field(foot, " PAGE ", "1")
    foot.add_run(" 页 共 ")
    _add_field(foot, " NUMPAGES ", "1")
    foot.add_run(" 页")


# --------------------------------------------------------------------------
# 样式
# --------------------------------------------------------------------------


def _define_styles(document) -> None:
    normal = document.styles["Normal"]
    _set_style_font(normal, east_asia=FONT_SONG, size_pt=10.5)

    for name, size, ea in (("Heading 1", 16, FONT_HEI), ("Heading 2", 14, FONT_HEI), ("Heading 3", 12, FONT_HEI)):
        style = document.styles[name]
        _set_style_font(style, east_asia=ea, size_pt=size, bold=True)
        style.font.color.rgb = None
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    # 每章另起一页。写在样式上而不是逐章插分页符：模板作者加一章就自动生效，
    # 生成器也不必知道哪儿该断页。已经在页首的标题不会因此多出一页空白——
    # Word 对页首段落的 pageBreakBefore 不再断一次。
    document.styles["Heading 1"].paragraph_format.page_break_before = True

    def add(name: str, **kwargs):
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = document.styles["Normal"]
        _set_style_font(style, **kwargs)
        return style

    # 正文宋体小四（12pt）。表格另有自己的样式，仍是五号。
    body = add(STYLE_BODY, east_asia=FONT_SONG, size_pt=12)
    body.paragraph_format.line_spacing = 1.5
    body.paragraph_format.space_after = Pt(0)
    _set_first_line_indent(body, 2)

    caption = add(STYLE_TABLE_CAPTION, east_asia=FONT_HEI, size_pt=10.5, bold=True)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(6)
    caption.paragraph_format.space_after = Pt(3)
    caption.paragraph_format.keep_with_next = True

    # 表头与表格文字都用宋体五号（10.5pt），表头加粗；一律水平居中。
    # 表格里没有首行缩进，行距按单倍，否则十列的病害表会被撑得很高。
    header_cell = add(STYLE_TABLE_HEADER, east_asia=FONT_SONG, size_pt=10.5, bold=True)
    header_cell.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_cell.paragraph_format.space_before = Pt(0)
    header_cell.paragraph_format.space_after = Pt(0)
    header_cell.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    cell = add(STYLE_TABLE_CELL, east_asia=FONT_SONG, size_pt=10.5)
    cell.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell.paragraph_format.space_before = Pt(0)
    cell.paragraph_format.space_after = Pt(0)
    cell.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    # 附录的两张卡片用小五（9pt）：格子比正文表格密得多，五号排不下就到处折行。
    card_header = add(STYLE_CARD_HEADER, east_asia=FONT_SONG, size_pt=9, bold=True)
    card_header.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    card_header.paragraph_format.space_before = Pt(0)
    card_header.paragraph_format.space_after = Pt(0)
    card_header.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    card_cell = add(STYLE_CARD_CELL, east_asia=FONT_SONG, size_pt=9)
    card_cell.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    card_cell.paragraph_format.space_before = Pt(0)
    card_cell.paragraph_format.space_after = Pt(0)
    card_cell.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    # 「A 桥梁所处行政区划代码」这类分段行整行合并，靠左排——居中会看着像标题。
    card_band = add(STYLE_CARD_BAND, east_asia=FONT_SONG, size_pt=9, bold=True)
    card_band.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    card_band.paragraph_format.space_before = Pt(0)
    card_band.paragraph_format.space_after = Pt(0)
    card_band.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    photo = add(STYLE_PHOTO, east_asia=FONT_SONG, size_pt=10.5)
    photo.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    photo.paragraph_format.space_after = Pt(0)
    # 图片与图题不可拆分：图片段落必须与下一段（图题）同页（设计 §11.5）。
    photo.paragraph_format.keep_with_next = True

    photo_caption = add(STYLE_PHOTO_CAPTION, east_asia=FONT_SONG, size_pt=9)
    photo_caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    photo_caption.paragraph_format.space_after = Pt(6)

    cover_title = add(STYLE_COVER_TITLE, east_asia=FONT_HEI, size_pt=26, bold=True)
    cover_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_title.paragraph_format.space_after = Pt(18)

    cover_field = add(STYLE_COVER_FIELD, east_asia=FONT_KAI, size_pt=14)
    cover_field.paragraph_format.space_after = Pt(10)
    cover_field.paragraph_format.left_indent = Cm(4.0)

    notice = add(STYLE_NOTICE, east_asia=FONT_SONG, size_pt=10.5)
    notice.paragraph_format.line_spacing = 1.5
    notice.paragraph_format.space_after = Pt(4)

    _add_table_styles(document)


def _add_table_styles(document) -> None:
    """两个表格样式：病害/人员/设备表带框线，照片版式不带。

    框线是样式的事，不是生成器的事——Docx Builder 只按名字取样式，改版式改这里
    （设计 §18）。python-docx 没有设置表格边框的 API，只能拼 XML。
    """
    table = document.styles.add_style(STYLE_TABLE, WD_STYLE_TYPE.TABLE)
    table.element.append(
        parse_xml(
            f"<w:tblPr {nsdecls('w')}>"
            f"<w:tblBorders>"
            + "".join(
                f'<w:{edge} w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
                for edge in ("top", "left", "bottom", "right", "insideH", "insideV")
            )
            + "</w:tblBorders>"
            '<w:tblCellMar>'
            '<w:top w:w="28" w:type="dxa"/><w:bottom w:w="28" w:type="dxa"/>'
            '<w:left w:w="57" w:type="dxa"/><w:right w:w="57" w:type="dxa"/>'
            "</w:tblCellMar>"
            "</w:tblPr>"
        )
    )

    layout = document.styles.add_style(STYLE_PHOTO_LAYOUT, WD_STYLE_TYPE.TABLE)
    layout.element.append(
        parse_xml(
            f"<w:tblPr {nsdecls('w')}>"
            "<w:tblBorders>"
            + "".join(
                f'<w:{edge} w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
                for edge in ("top", "left", "bottom", "right", "insideH", "insideV")
            )
            + "</w:tblBorders>"
            '<w:tblCellMar>'
            '<w:top w:w="0" w:type="dxa"/><w:bottom w:w="0" w:type="dxa"/>'
            '<w:left w:w="57" w:type="dxa"/><w:right w:w="57" w:type="dxa"/>'
            "</w:tblCellMar>"
            "</w:tblPr>"
        )
    )


# --------------------------------------------------------------------------
# 各页
# --------------------------------------------------------------------------


def _cover(document) -> None:
    for _ in range(4):
        document.add_paragraph("", style=STYLE_BODY)
    document.add_paragraph(
        "{{inspection_year}}年{{administrative_region}}{{route_code}}{{route_name}}",
        style=STYLE_COVER_TITLE,
    )
    document.add_paragraph("{{bridge_name}}定期检测报告", style=STYLE_COVER_TITLE)
    for _ in range(6):
        document.add_paragraph("", style=STYLE_BODY)
    for line in (
        "委托单位：",
        "检测单位：{{inspection_org}}",
        "检测日期：{{inspection_date}}",
        "报告日期：{{report_date}}",
        "报告编号：{{report_no}}",
    ):
        document.add_paragraph(line, style=STYLE_COVER_FIELD)


def _declaration(document) -> None:
    document.add_paragraph("注意事项", style=STYLE_COVER_TITLE)
    for clause in NOTICE_CLAUSES:
        document.add_paragraph(clause, style=STYLE_NOTICE)
    document.add_paragraph("", style=STYLE_BODY)
    document.add_paragraph("检测单位：{{inspection_org}}", style=STYLE_NOTICE)
    for line in ("地址：", "邮政编码：", "联系人：", "电话：", "传真：", "电子邮箱："):
        document.add_paragraph(line, style=STYLE_NOTICE)


def _signature_page(document) -> None:
    document.add_paragraph("工程（产品）名称：{{bridge_name}}", style=STYLE_NOTICE)
    document.add_paragraph("签字表", style=STYLE_COVER_TITLE)
    document.add_paragraph(anchor_text("PERSONNEL_TABLE"), style=STYLE_BODY)
    document.add_paragraph("", style=STYLE_BODY)
    document.add_paragraph(
        "本页签字栏由报告签发人手写签署，系统不自动插入电子签名。", style=STYLE_NOTICE
    )


def _table_of_contents(document) -> None:
    document.add_paragraph("目　录", style=STYLE_COVER_TITLE)
    paragraph = document.add_paragraph("", style=STYLE_BODY)
    _add_field(paragraph, ' TOC \\o "1-3" \\h \\z \\u ', "右键“更新域”生成目录。")


def _chapter_one(document) -> None:
    document.add_heading("1 项目概况", level=1)

    document.add_heading("1.1 桥梁概况", level=2)
    document.add_paragraph(anchor_text("BRIDGE_PROFILE"), style=STYLE_BODY)

    document.add_heading("1.2 检测目的", level=2)
    document.add_paragraph(
        "按照公路桥梁养护规范要求，对该桥进行定期检测，掌握其技术状况，"
        "为养护决策提供依据。",
        style=STYLE_BODY,
    )

    document.add_heading("1.3 检测依据", level=2)
    for item in (
        "《公路桥涵养护规范》（JTG 5120—2021）",
        "《公路桥梁技术状况评定标准》（JTG/T H21—2011）",
        "《公路技术状况评定标准》（JTG 5210—2018）",
    ):
        document.add_paragraph(item, style=STYLE_BODY)

    document.add_heading("1.4 桥梁养护历史", level=2)
    document.add_heading("1.4.1 历年检测情况", level=3)
    document.add_paragraph("", style=STYLE_BODY)
    document.add_heading("1.4.2 维修加固情况", level=3)
    document.add_paragraph("", style=STYLE_BODY)

    document.add_heading("1.5 构件编号", level=2)
    document.add_heading("1.5.1 方位描述规则", level=3)
    document.add_paragraph(
        "以桩号增大方向为前进方向，左右侧按前进方向确定。", style=STYLE_BODY
    )
    document.add_heading("1.5.2 构件编号规则", level=3)
    document.add_paragraph(
        "构件编号采用「孔号-构件序号#构件类型」的完整记号，与构件台账保持一致。",
        style=STYLE_BODY,
    )

    document.add_heading("1.6 检测人员与设备", level=2)
    document.add_paragraph(anchor_text("EQUIPMENT_LIST"), style=STYLE_BODY)


def _chapter_two(document) -> None:
    document.add_heading("2 桥梁缺损状况检查", level=1)
    for number, name, part in STRUCTURE_SECTIONS:
        document.add_heading(f"{number} {name}", level=2)

        document.add_heading(f"{number}.1 {name}检查", level=3)
        document.add_paragraph(anchor_text("DEFECT_TABLES", part), style=STYLE_BODY)
        document.add_paragraph(anchor_text("DEFECT_PHOTOS", part), style=STYLE_BODY)

        document.add_heading(f"{number}.2 与所选历史检查结果对比", level=3)
        document.add_paragraph(anchor_text("PREVIOUS_COMPARISON", part), style=STYLE_BODY)

        document.add_heading(f"{number}.3 病害成因分析", level=3)
        document.add_paragraph("", style=STYLE_BODY)


def _chapter_three(document) -> None:
    """材质状况检测数据不在本次范围内（设计 §3.5），保留标题和版式，正文留空。"""
    document.add_heading("3 桥梁材质状况检测评定", level=1)
    for number, title in (
        ("3.1", "回弹法测试构件混凝土强度"),
        ("3.2", "混凝土碳化状况检测评定"),
        ("3.3", "钢筋锈蚀电位检测评定"),
    ):
        document.add_heading(f"{number} {title}", level=2)
        document.add_paragraph("", style=STYLE_BODY)


def _chapter_four(document) -> None:
    """第 4 章按正式报告的分节：4.1 综合评定（权重分配 + 技术状况等级）、
    4.2 单项控制指标、4.3 等级综合评定。"""
    document.add_heading("4 全桥技术状况综合评定", level=1)

    document.add_heading("4.1 桥梁技术状况综合评定", level=2)
    document.add_heading("4.1.1 部件权重分配", level=3)
    document.add_paragraph(anchor_text("COMPONENT_WEIGHTS"), style=STYLE_BODY)
    document.add_heading("4.1.2 桥梁技术状况等级", level=3)
    document.add_paragraph(anchor_text("ASSESSMENT_RESULT"), style=STYLE_BODY)

    document.add_heading("4.2 桥梁技术状况等级单项控制指标", level=2)
    document.add_paragraph(anchor_text("CONTROL_INDICATOR"), style=STYLE_BODY)

    document.add_heading("4.3 桥梁技术状况等级综合评定", level=2)
    document.add_paragraph(anchor_text("OVERALL_ASSESSMENT"), style=STYLE_BODY)


def _chapter_five(document) -> None:
    document.add_heading("5 检测结论与养护建议", level=1)
    document.add_paragraph(anchor_text("CONCLUSION"), style=STYLE_BODY)


def _appendix(document) -> None:
    document.add_heading("6 附录", level=1)
    document.add_heading("附录1 桥梁技术状况评定表", level=2)
    document.add_paragraph(anchor_text("ASSESSMENT_APPENDIX"), style=STYLE_BODY)
    document.add_heading("附录2 桥梁基本状况卡片", level=2)
    document.add_paragraph("", style=STYLE_BODY)


# --------------------------------------------------------------------------
# 组装
# --------------------------------------------------------------------------


def build_standard_template(path: Path) -> Path:
    document = Document()
    _define_styles(document)

    _portrait(document.sections[0])
    _blank_header_footer(document.sections[0])
    _cover(document)

    for builder in (_declaration, _signature_page, _table_of_contents):
        section = document.add_section(WD_SECTION.NEW_PAGE)
        _portrait(section)
        _blank_header_footer(section)
        builder(document)

    body_section = document.add_section(WD_SECTION.NEW_PAGE)
    _portrait(body_section)
    _running_header_footer(body_section)
    # 正文才是第 1 页：封面、声明、签字、目录不参与正文页码。
    _set_page_numbering_restart(body_section, 1)
    for builder in (
        _chapter_one,
        _chapter_two,
        _chapter_three,
        _chapter_four,
        _chapter_five,
    ):
        builder(document)

    # 附录的评定表和基本状况卡片都很宽，与正式报告一致走横向节。
    appendix_section = document.add_section(WD_SECTION.NEW_PAGE)
    _landscape(appendix_section)
    _running_header_footer(appendix_section)
    _appendix(document)

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return path


def build_template_config() -> dict:
    """随模板一起发布的配置，即管理员登记模板时写进 contract_config_json 的内容。"""
    return {
        "contract_type": CONTRACT_PERIODIC_INSPECTION_V1,
        "table_number_formats": dict(TEMPLATE_NUMBER_FORMATS),
        "required_personnel_roles": list(TEMPLATE_REQUIRED_ROLES),
    }


def build(output_dir: Path) -> tuple[Path, Path]:
    docx_path = build_standard_template(output_dir / f"{TEMPLATE_BASENAME}.docx")
    config_path = output_dir / "template.json"
    config_path.write_text(
        json.dumps(build_template_config(), ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return docx_path, config_path


def main() -> None:
    if len(sys.argv) > 2:
        print(__doc__)
        raise SystemExit(2)
    output_dir = Path(sys.argv[1]) if len(sys.argv) == 2 else DEFAULT_OUTPUT_DIR

    # 模板已经转为手工维护：产物 .docx 才是真源，这个脚本只用来生成初版。
    # 已经存在就拒绝覆盖——它现在是人在 Word 里改出来的，跑一次就全没了。
    existing = output_dir / f"{TEMPLATE_BASENAME}.docx"
    if existing.exists():
        print(
            f"拒绝覆盖已存在的模板：{existing}\n"
            "\n"
            "模板现在由人在 Word 里维护，这个脚本只负责生成初版（设计 §8）。\n"
            "确实要重新生成初版时，请指定一个空目录：\n"
            f"    python {Path(__file__).name} <空目录>\n"
            "再手工比对、合并回正式模板。"
        )
        raise SystemExit(1)

    docx_path, config_path = build(output_dir)
    print(f"模板：{docx_path}")
    print(f"配置：{config_path}")


if __name__ == "__main__":
    main()
