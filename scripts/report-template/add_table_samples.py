"""往已有模板里补一个「表头样表区」（设计 §7.7）。

模板已转为手工维护，所以这个脚本只做加法：先备份，再在正文节末尾插入若干
「[[TABLE:内容块]] + 只有表头行的表」，其余内容一个字不动。已经有样表的内容块
直接跳过，可以重复跑。

插好之后你在 Word 里拖列宽、改表头文字，生成的报告就照着走；行数和合并仍由数据
决定。样表本身不会出现在报告里——生成时连标记带表一起删掉。

样表放在它对应的表格真正渲染的那一节里：正文表放正文节（版心 16cm），附表1 放
横向的附录节（24.62cm）。放错节的话你在 Word 里量到的宽度和实际排版对不上。

用法（在仓库根目录）：
    tools-python/.venv/Scripts/python.exe scripts/report-template/add_table_samples.py
"""

from __future__ import annotations

import shutil
import sys
from datetime import datetime
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO_ROOT / "tools-python"))

from docx import Document  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm  # noqa: E402

from bridge_report_tools.reports.contract import TABLE_COLUMNS, table_marker  # noqa: E402
from bridge_report_tools.reports.docx_scan import scan_table_samples  # noqa: E402
from bridge_report_tools.reports.styles import STYLE_BODY, STYLE_TABLE  # noqa: E402

TEMPLATE = (
    REPO_ROOT / "templates" / "report" / "periodic_inspection_v1"
    / "periodic-inspection-v1.docx"
)

#: 各内容块的初始列宽（cm）。只是个起点——你在 Word 里拖成什么样就是什么样。
#: 正文版心 16cm，附录横向节 24.62cm。
INITIAL_WIDTHS = {
    # 序号 部件名称 构件编号 病害位置 病害类型 病害特征 标度 病害扣分 构件评分 照片编号
    #
    # 「病害特征」和「照片编号」是两个吃宽度的列：前者常有二三十字的描述，后者
    # 一条病害挂三张照片就是「照片2.1-2、照片2.1-3、照片2.1-4」。给窄了，297 行的
    # 上部结构病害表会被撑高一倍——实测把病害特征压到 3.3cm，报告从 104 页涨到 184 页。
    "DEFECT_TABLES": [0.8, 1.5, 1.4, 1.5, 1.7, 4.4, 0.7, 0.9, 0.9, 2.2],
    # 部位 序号 名称 权重 重新分配后权重 构件数量 备注
    "COMPONENT_WEIGHTS": [1.8, 0.9, 4.2, 1.6, 2.6, 1.7, 3.2],
    # 结构 类别 评价部件 构件数量 构件评分 部件评分 结构评分 权重 等级 全桥评分 综合评级
    "ASSESSMENT_RESULT": [1.5, 0.8, 3.6, 1.1, 1.2, 1.4, 1.4, 1.4, 0.8, 1.4, 1.4],
    # 序号 姓名 单位 职称 资格证书编号 职责
    "PERSONNEL_TABLE": [1.2, 2.2, 4.4, 2.6, 3.4, 2.2],
    # 序号 设备名称 型号规格 资产编号 量程 精度 检定证书编号 检定有效期 用途
    "EQUIPMENT_LIST": [1.0, 2.6, 2.2, 1.8, 1.6, 1.4, 2.4, 1.8, 1.2],
    # 附表1 在横向的附录节，版心 24.62cm。
    "ASSESSMENT_APPENDIX": [2.4, 1.2, 5.0, 2.0, 3.4, 3.4, 3.2, 2.0],
}

#: 各内容块的样表放在第几节（0 起）。正文节是第 5 节，附录节是第 6 节。
BODY_SECTION_INDEX = 4
APPENDIX_SECTION_INDEX = 5
SECTION_OF_BLOCK = {
    "DEFECT_TABLES": BODY_SECTION_INDEX,
    "COMPONENT_WEIGHTS": BODY_SECTION_INDEX,
    "ASSESSMENT_RESULT": BODY_SECTION_INDEX,
    "PERSONNEL_TABLE": BODY_SECTION_INDEX,
    "EQUIPMENT_LIST": BODY_SECTION_INDEX,
    "ASSESSMENT_APPENDIX": APPENDIX_SECTION_INDEX,
}

#: 每张样表的说明，写在标记段的标记后面。整段生成时会被删掉，说明也就跟着走，
#: 不必另起一段——另起的段留在报告里就成了莫名其妙的一行字。
HINTS = {
    "DEFECT_TABLES": "病害检查表：拖列宽、改表头文字即可；行数由病害条数决定。",
    "COMPONENT_WEIGHTS": "表4.1-1 部件权重计算表：行数由桥型的规范部件清单决定。",
    "ASSESSMENT_RESULT": "表4.1-2 总体技术状况评定表：行数与合并跨度由构件评分分档决定。",
    "PERSONNEL_TABLE": "签字页人员表。",
    "EQUIPMENT_LIST": "检测设备表。",
    "ASSESSMENT_APPENDIX": "附表1（横向节，版心 24.62cm）。",
}


def section_break_paragraphs(document):
    """带分节符的段落，按文档顺序。第 i 个就是第 i 节的末尾。"""
    return [
        child
        for child in document.element.body.iterchildren()
        if child.tag == qn("w:p") and child.xpath("./w:pPr/w:sectPr")
    ]


def build_sample(document, block: str, widths_cm: list[float]):
    """造一个「标记段 + 表头样表」，返回这两个 XML 元素。"""
    marker = document.add_paragraph(
        f"{table_marker(block)}  {HINTS.get(block, '')}".rstrip(), style=STYLE_BODY
    )
    headers = TABLE_COLUMNS[block]
    table = document.add_table(1, len(headers), style=STYLE_TABLE)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, header in enumerate(headers):
        table.cell(0, index).text = header
        table.columns[index].width = Cm(widths_cm[index])
        table.rows[0].cells[index].width = Cm(widths_cm[index])
    return marker._p, table._tbl


def main() -> None:
    target_path = Path(sys.argv[1]) if len(sys.argv) > 1 else TEMPLATE
    if not target_path.is_file():
        print(f"找不到模板：{target_path}")
        raise SystemExit(1)

    document = Document(str(target_path))
    existing = {sample.block for sample in scan_table_samples(document)}
    todo = [block for block in TABLE_COLUMNS if block not in existing]
    if not todo:
        print("每个内容块都已经有样表了，没有要补的。")
        return

    breaks = section_break_paragraphs(document)
    if len(breaks) <= BODY_SECTION_INDEX:
        print(
            f"模板的分节结构与预期不符（找到 {len(breaks)} 个分节符，"
            f"至少需要 {BODY_SECTION_INDEX + 1} 个）。请手工插入样表。"
        )
        raise SystemExit(1)

    backup = target_path.with_suffix(f".{datetime.now():%Y%m%d-%H%M%S}.bak.docx")
    shutil.copy2(target_path, backup)

    # 正文节的样表插在正文节的分节符之前；附录节的留在文档末尾，
    # python-docx 的 add_* 本来就往那儿追加。
    body_anchor = breaks[BODY_SECTION_INDEX]

    def place(element, section_index: int) -> None:
        if section_index == BODY_SECTION_INDEX:
            body_anchor.addprevious(element)

    for block in todo:
        section_index = SECTION_OF_BLOCK[block]
        marker, table = build_sample(document, block, INITIAL_WIDTHS[block])
        place(marker, section_index)
        place(table, section_index)
        # 表格后面补一个空段落：两张表挨着会被 Word 合并成一张。
        place(document.add_paragraph("", style=STYLE_BODY)._p, section_index)

    document.save(str(target_path))
    print(f"备份：{backup}")
    print(f"已补样表：{'、'.join(todo)}")
    print(f"模板：{target_path}")


if __name__ == "__main__":
    main()
