import hashlib,importlib.util,json,sys
from pathlib import Path
from docx import Document
from docx.shared import Inches
from PIL import Image,ImageDraw
ROOT=Path(__file__).resolve().parents[2]
environment=json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
out=Path(environment["run_directory"])/"samples";out.mkdir(exist_ok=True)
spec=importlib.util.spec_from_file_location("fixtures",ROOT/"tools-python/tests/importers/docx_fixtures.py")
fixtures=importlib.util.module_from_spec(spec);spec.loader.exec_module(fixtures)
doc=Document();doc.add_paragraph("验收小样桥 定期检测报告")
fixtures.add_liaoning_defect_table(doc)
for row in doc.tables[0].rows:
 for cell in row.cells: cell.text=cell.text.replace("2-1#板","2-1#梁")
for i,caption in [(1,"蜂窝、麻面"),(2,"剥落、掉角")]:
 p=out/f"photo-{i}.png";im=Image.new("RGB",(400,240),(220,220,220));d=ImageDraw.Draw(im);d.rectangle((50,50,200,150),fill=(100+i*30,80,70));d.text((15,15),f"TEST PHOTO {i}",fill="black");im.save(p)
 doc.add_picture(str(p),width=Inches(2));doc.add_paragraph(f"照片2.1-{i} 2-1#梁{caption}")
doc.save(out/"small.docx")
manifest={"source":str(out/"small.docx"),"expected_defects":2,"expected_photos":2,"expected_component":"2-1#梁","expected_inventory_count":41,"photo_captions":["2.1-1","2.1-2"],"input_sha256":hashlib.sha256((out/"small.docx").read_bytes()).hexdigest(),"note":"合成辽宁表头；来源评分不作为系统评分预期。"}
(out/"expected.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding="utf-8")
print("Small source prepared: 2 defects, 2 photos, component 2-1#梁.")
