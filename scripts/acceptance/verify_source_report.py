"""Check the downloaded 2026 source-database report against formal facts and DOCX content."""
from __future__ import annotations
import argparse, hashlib, json, re
from collections import Counter
from pathlib import Path
from zipfile import ZipFile
from docx import Document

def read(path: Path):
    return json.loads(path.read_text(encoding="utf-8-sig"))
def text(value):
    return "—" if value is None or value == "" else str(value)
def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("manifest",type=Path)
    args=ap.parse_args()
    manifest=read(args.manifest)
    if manifest["database"]!="bridge_report_acceptance" or manifest["ports"]["postgres"]!=25432:
        raise SystemExit("not the isolated source acceptance database")
    root=Path(manifest["run_directory"]);log=Path(manifest["log_directory"])
    state=read(root/"scenario.json")
    entry=state["sourceReport"]
    job=entry["latestJob"]
    year=state["sourceDatabase"]["years"]["2026"]["year"]["id"]
    assert job["inspection_year_id"]==year and job["status"]=="ready"
    docx=Path(entry["download"])
    context=read(root/"report-jobs"/job["id"]/"context.json")
    formal=read(log/"source-confirm-database.json")
    inventory=read(log/"source-assessment-raw.json")["inventory"]
    number_by_id={c["bridge_component_id"]:c["component_number"] for c in inventory}
    errors=[]
    def check(name,actual,expected):
        if actual!=expected:errors.append({"check":name,"actual":str(actual)[:180],"expected":str(expected)[:180]})
    check("context_year",context["inspection_year_id"],year)
    check("formal_year",formal["year"]["id"],year)
    check("context_score",context["assessment"]["overall_score"],81.10733)
    check("context_grade",context["assessment"]["overall_grade"],"2类")
    check("context_source_count",context["overall_comparison"]["current_source_defect_count"],381)
    source={o["id"]:o for o in formal["observations"]}
    check("source_observation_count",len(source),430)
    context_rows=[r for p in context["parts"] for r in p["defect_rows"]]
    check("context_row_count",len(context_rows),430)
    check("context_observation_ids",set(r["observation_id"] for r in context_rows),set(source))
    report_photo_by_number={p["report_number"]:p for part in context["parts"] for p in part["photos"]}
    check("context_photo_count",len(report_photo_by_number),455)
    source_photo_by_id={p["id"]:p for o in source.values() for p in o["photos"]}
    check("source_photo_ids",set(p["photo_id"] for p in report_photo_by_number.values()),set(source_photo_by_id))
    for row in context_rows:
        obs=source[row["observation_id"]]
        key=row["observation_id"]
        check("component:"+key,row["component_number"],number_by_id[obs["bridge_component_id"]])
        for report_key,source_key in (("defect_location","defect_location"),("defect_type","defect_type"),
                                      ("description","defect_description_raw"),("scale","scale")):
            check(report_key+":"+key,row[report_key],obs[source_key])
        actual_photo_ids={report_photo_by_number[number]["photo_id"] for number in row["photo_numbers"]}
        check("photo_owner:"+key,actual_photo_ids,{p["id"] for p in obs["photos"]})
    document=Document(str(docx))
    check("docx_section_count",len(document.sections),7)
    check("docx_image_shapes",len(document.inline_shapes),455)
    for part,table_index,photo_table_index in zip(context["parts"],(2,4,6),(3,5,7)):
        table=document.tables[table_index]
        check("table_rows:"+part["part_code"],len(table.rows),len(part["defect_rows"])+1)
        for expected,actual in zip(part["defect_rows"],table.rows[1:]):
            fields=(expected["row_number"],expected["part_name"],expected["component_number"],
                    expected["defect_location"],expected["defect_type"],expected["description"],
                    expected["scale"],expected["deduction"],expected["component_score"],
                    "、".join(expected["photo_numbers"]) if expected["photo_numbers"] else None)
            displayed=[c.text for c in actual.cells]
            check("docx_row:"+expected["observation_id"],displayed,[text(v) for v in fields])
        photo_table=document.tables[photo_table_index]
        captions=[cell.text for row in photo_table.rows for cell in row.cells if cell.text]
        check("photo_captions:"+part["part_code"],captions,[p["caption"] for p in part["photos"]])
    with ZipFile(docx) as z:
        check("zip_integrity",z.testzip(),None)
        media=[n for n in z.namelist() if n.startswith("word/media/") and not n.endswith("/")]
    check("embedded_media",len(media),455)
    visible="\n".join([p.text for p in document.paragraphs]+
        [c.text for t in document.tables for row in t.rows for c in row.cells])
    check("unresolved_template_marker",bool(re.search(r"\{\{|\[\[BLOCK:",visible)),False)
    check("report_score_visible","81.1 分" in visible or "81.1" in visible,True)
    check("grade_visible","2类" in visible,True)
    check("no_previous_visible","暂无可供对比的历史正式检查记录" in visible,True)
    meta={k:context["scalars"].get(k) for k in ("inspection_date","report_no","project_name","inspection_org")}
    pdf_report=read(log/("source-report-render-fixed.json" if entry.get("previousJob") else "source-report-render.json"))
    check("pdf_pages",pdf_report["pages"],job["progress"]["page_count"])
    check("pdf_blank_pages",pdf_report["pages_under_20_chars"],[])
    if entry.get("previousJob"):
        expected_meta={"inspection_date":"2026-03-18","report_no":"ACCEPT-REPORT-2026",
                       "project_name":"百股大桥2026年度定期检测（隔离验收）",
                       "inspection_org":"验收检测机构"}
        check("manual_cover_metadata",meta,expected_meta)
        for value in ("ACCEPT-REPORT-2026", "验收检测机构", "2026-03-18"):
            check("visible_cover:"+value,value in visible,True)
        for physical,printed in ((97,92),(98,93),(99,94),(100,95)):
            check("appendix_page_number:"+str(physical),
                  pdf_report["selected"][str(physical)]["footer"],
                  [f"第 {printed} 页 共 100 页"])
        check("template_snapshot_checksum",job["template_checksum"],state["template"]["file_checksum"])
    report={"passed":not errors,"job_id":job["id"],"year_id":year,
            "docx_sha256":hashlib.sha256(docx.read_bytes()).hexdigest(),
            "docx_bytes":docx.stat().st_size,"pdf_pages":pdf_report["pages"],
            "formal_observations":len(source),"report_rows":len(context_rows),
            "formal_photos":len(source_photo_by_id),"report_photo_captions":len(report_photo_by_number),
            "docx_inline_images":len(document.inline_shapes),"docx_media_parts":len(media),
            "part_rows":{p["part_code"]:len(p["defect_rows"]) for p in context["parts"]},
            "part_photos":{p["part_code"]:len(p["photos"]) for p in context["parts"]},
            "report_score_printed":"81.1","system_year_score":"81.11",
            "cover_metadata":meta,"error_count":len(errors),"errors":errors[:100]}
    (log/"source-report-verification.json").write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding="utf-8")
    print(json.dumps(report,ensure_ascii=False,indent=2))
    if errors:raise SystemExit(1)
if __name__=="__main__":main()

